"""Protocol-parameterised endpoint tests for site/assets/template/run_analysis.py.

Run: python3 -I -m unittest discover -s tests -p test_analysis_protocol.py
Requires the pinned analysis dependencies (site/assets/template/requirements.txt:
numpy, pandas, openpyxl, matplotlib). Synthetic fixtures: tests/fixtures/analysis/;
further synthetic scenarios are built per test (fictitious IDs, no real patients).
"""
from __future__ import annotations

import contextlib
import copy
import csv
import hashlib
import io
import json
import platform
import re
import shutil
import sys
import tempfile
import unittest
from datetime import date, timedelta
from fractions import Fraction
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "site" / "assets" / "template"
FIXTURES = ROOT / "tests" / "fixtures" / "analysis"
DATA = FIXTURES / "data"
PROTOCOL = FIXTURES / "protocol_confirmed.json"
sys.path.insert(0, str(TEMPLATE))
try:
    import pandas as pd  # noqa: E402
    import run_analysis as ra  # noqa: E402
    import merge_centers as mc  # noqa: E402
except ImportError as exc:  # fail loudly: a silent skip would hide the endpoint logic from CI
    raise ImportError(f"test_analysis_protocol needs the analysis dependencies in {TEMPLATE/'requirements.txt'}: {exc}") from exc

P1 = "00000000-0000-4000-8000-000000000001"
P2 = "00000000-0000-4000-8000-000000000002"
SP = "00000000-0000-4000-8000-0000000000aa"  # synthetic scenarios built inside tests
B0 = date(2020, 1, 1)
BASE_HEADER = ra.KEY + ["id", "module", "sex", "birth_year", "baseline_date", "baseline_scr", "baseline_upcr", "baseline_upcr_unit"]
VISIT_HEADER = ra.KEY + ["id", "module", "visit_date", "sbp", "dbp", "scr_umol_l", "upcr", "egfr", "egfr_formula_version"]
EVENT_HEADER = ra.KEY + ["id", "module", "event_type", "event_date", "confirmed", "source"]


def day(n):
    return (B0 + timedelta(days=n)).isoformat()


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, header, rows):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with Path(path).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header, lineterminator="\n")
        w.writeheader()
        w.writerows([{k: r.get(k, "") for k in header} for r in rows])


def make_data(folder, patients, events=(), baseline_cols=()):
    """patients: dicts with code, center, module, bupcr, baseline_date, extra (baseline
    columns) and visits [(day, egfr, upcr[, formula_version[, scr_umol_l]])]; None = blank."""
    base_rows, visit_rows, event_rows = [], [], []
    for p in patients:
        center, module = p.get("center", "A"), p.get("module", "IGAN")
        base_rows.append({"project_id": SP, "center_code": center, "patient_code": p["code"], "id": f"b{center}{p['code']}",
                          "module": module, "sex": "M", "birth_year": "1970", "baseline_date": p.get("baseline_date", day(0)),
                          "baseline_scr": "80", "baseline_upcr": p.get("bupcr", "2000"), "baseline_upcr_unit": "mg/g", **p.get("extra", {})})
        for j, v in enumerate(p.get("visits", []), start=1):
            dd, egfr, upcr = v[:3]
            visit_rows.append({"project_id": SP, "center_code": center, "patient_code": p["code"], "id": f"v{center}{p['code']}-{j}",
                               "module": module, "visit_date": day(dd), "sbp": "120", "dbp": "80",
                               "scr_umol_l": v[4] if len(v) > 4 else "", "upcr": "" if upcr is None else str(upcr),
                               "egfr": "" if egfr is None else str(egfr), "egfr_formula_version": v[3] if len(v) > 3 else "manual"})
    for i, e in enumerate(events, start=1):
        center = e.get("center", "A")
        module = next(p.get("module", "IGAN") for p in patients if p["code"] == e["code"] and p.get("center", "A") == center)
        event_rows.append({"project_id": SP, "center_code": center, "patient_code": e["code"], "id": f"e{i:03d}", "module": module,
                           "event_type": e["type"], "event_date": day(e["day"]), "confirmed": e.get("confirmed", "true"), "source": "staff"})
    write_csv(Path(folder) / "patients_baseline.csv", BASE_HEADER + list(baseline_cols), base_rows)
    write_csv(Path(folder) / "visits_long.csv", VISIT_HEADER, visit_rows)
    write_csv(Path(folder) / "events_long.csv", EVENT_HEADER, event_rows)
    return Path(folder)


def rows_for(rows, code, endpoint=None, center="A"):
    out = [r for r in rows if r["center_code"] == center and r["patient_code"] == code and (endpoint is None or r["endpoint"] == endpoint)]
    return out


class AnalysisProtocolTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="analysis-protocol-test-"))
        cls.protocol_out = cls.tmp / "protocol-run"
        cls.protocol_log = ra.analyze(DATA, cls.protocol_out, PROTOCOL)
        cls.endpoints = read_csv(cls.protocol_out / "endpoints_per_protocol.csv")
        cls.candidate_out = cls.tmp / "candidate-run"
        cls.candidate_log = ra.analyze(DATA, cls.candidate_out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def ep(self, project, center, code, endpoint, rows=None):
        rows = [r for r in (self.endpoints if rows is None else rows) if (r["project_id"], r["center_code"], r["patient_code"], r["endpoint"]) == (project, center, code, endpoint)]
        self.assertEqual(len(rows), 1, (project, center, code, endpoint))
        return rows[0]

    def newdir(self):
        return Path(tempfile.mkdtemp(dir=self.tmp))

    def variant_protocol(self, mutate, name="variant.json", base=None):
        raw = json.loads((base or PROTOCOL).read_text(encoding="utf-8"))
        mutate(raw)
        path = self.newdir() / name
        path.write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")
        return path

    def synth_protocol(self, mutate=None):
        """Fixture protocol with no censoring event types (synthetic data has no events
        unless a test adds them and sets censoring.event_types itself)."""
        def m(raw):
            raw["censoring"]["event_types"] = []
            if mutate:
                mutate(raw)
        return self.variant_protocol(m)

    def synth(self, patients, events=(), baseline_cols=()):
        return make_data(self.newdir() / "data", patients, events, baseline_cols)

    def run_protocol(self, data, protocol, **kw):
        out = self.newdir() / "out"
        log = ra.analyze(data, out, protocol, **kw)
        return log, read_csv(out / "endpoints_per_protocol.csv"), read_csv(out / "qc_issues.csv"), out

    def variant_data(self, mutate_baseline=None, mutate_visits=None, dictionary=None, extra_tables=None):
        out = self.newdir() / "data"
        shutil.copytree(DATA, out)
        for name, fn in [("patients_baseline.csv", mutate_baseline), ("visits_long.csv", mutate_visits)]:
            if fn is None:
                continue
            rows = read_csv(out / name)
            header, rows = fn(list(rows[0].keys()), rows)
            write_csv(out / name, header, rows)
        for name, (header, rows) in (extra_tables or {}).items():
            existing = read_csv(out / name) if (out / name).exists() else []
            write_csv(out / name, header, existing + rows)
        if dictionary is not None:
            (out / "data_dictionary.json").write_text(json.dumps(dictionary), encoding="utf-8")
        return out

    # --- sustained eGFR decline -------------------------------------------------
    def test_transient_decline_then_recovery_is_not_sustained(self):
        r = self.ep(P1, "A", "001", "egfr_decline_40pct_sustained")
        self.assertEqual((r["status"], r["status_detail"]), ("not_met", "transient_not_confirmed_recovered"))
        self.assertEqual(r["event_date"], "")

    def test_sustained_decline_is_determined_with_event_and_confirmation_dates(self):
        r = self.ep(P1, "B", "001", "egfr_decline_40pct_sustained")
        self.assertEqual(r["status"], "met_per_protocol")
        self.assertEqual((r["event_date"], r["confirmation_date"]), ("2020-03-31", "2020-05-30"))
        self.assertEqual(r["determination"], "已按方案 TEST-SYNTHETIC@1.0 判定")
        self.assertEqual((r["event_basis"], r["n_unknown_measurements"], r["egfr_source_counts"]), ("egfr_measurements", "0", "manual:3"))
        self.assertEqual(self.ep(P1, "B", "001", "egfr_decline_57pct_sustained")["status"], "not_met")

    def test_interval_shorter_than_protocol_minimum_is_not_determined(self):
        r = self.ep(P1, "A", "002", "egfr_decline_40pct_sustained")
        self.assertEqual((r["status"], r["status_detail"]), ("not_met", "confirmation_interval_too_short"))

    def test_require_last_below_rejects_later_recovery(self):
        self.assertEqual(self.ep(P1, "A", "008", "egfr_decline_40pct_sustained")["status"], "met_per_protocol")
        path = self.variant_protocol(lambda p: p["egfr_decline"]["confirmation"].update(require_last_below=True))
        ra.analyze(DATA, self.tmp / "last-below", path)
        rows = read_csv(self.tmp / "last-below" / "endpoints_per_protocol.csv")
        r = [x for x in rows if (x["project_id"], x["center_code"], x["patient_code"], x["endpoint"]) == (P1, "A", "008", "egfr_decline_40pct_sustained")][0]
        self.assertEqual((r["status"], r["status_detail"]), ("not_met", "met_then_not_sustained_to_last_measurement"))

    def test_evaluate_sustained_unit_cases(self):
        T = pd.Timestamp
        conf = {"min_measurements": 2, "min_interval_days": 28, "require_last_below": False}
        seq = [(T("2020-01-10"), 9, True), (T("2020-02-10"), 40, True)]
        self.assertEqual(ra.evaluate_sustained(seq, conf)["status"], "met_per_protocol")
        seq = [(T("2020-01-10"), 9, True), (T("2020-01-20"), 19, True), (T("2020-02-20"), 50, False)]
        self.assertEqual(ra.evaluate_sustained(seq, conf)["status_detail"], "confirmation_interval_too_short")
        self.assertEqual(ra.evaluate_sustained([], conf)["status"], "not_evaluable")
        self.assertEqual(ra.evaluate_sustained([(T("2020-01-10"), 9, True)], conf)["status_detail"], "met_at_last_measurement_awaiting_confirmation")

    def test_evaluate_sustained_unknown_measurement_interrupts_and_is_never_qualifying(self):
        T = pd.Timestamp
        conf = {"min_measurements": 2, "min_interval_days": 28, "require_last_below": False}
        # An unknown between two qualifying values is not bridged into a run.
        seq = [(T("2020-04-01"), 91, True), (T("2020-05-01"), 121, None), (T("2020-06-01"), 152, True)]
        self.assertEqual(ra.evaluate_sustained(seq, conf), {"status": "not_met", "status_detail": "unknown_measurement_prevents_confirmation"})
        # An unknown alone never confirms: one known + unknown does not count as two.
        seq = [(T("2020-04-01"), 91, True), (T("2020-05-01"), 121, None)]
        self.assertEqual(ra.evaluate_sustained(seq, conf)["status"], "not_met")
        # When the unknown could not have changed the result, the ordinary detail stays.
        seq = [(T("2020-04-01"), 91, True), (T("2020-04-05"), 95, None), (T("2020-04-10"), 100, False)]
        self.assertEqual(ra.evaluate_sustained(seq, conf)["status_detail"], "not_confirmed_interrupted_by_unknown_measurement")
        self.assertEqual(ra.evaluate_sustained([(T("2020-04-01"), 91, None)], conf), {"status": "not_evaluable", "status_detail": "no_usable_followup_measurements"})
        # A confirmed run before an unknown is still met.
        seq = [(T("2020-04-01"), 91, True), (T("2020-05-01"), 121, True), (T("2020-06-01"), 152, None)]
        self.assertEqual(ra.evaluate_sustained(seq, conf)["status"], "met_per_protocol")
        # require_last_below: an unknown last measurement blocks confirmation.
        last = dict(conf, require_last_below=True)
        self.assertEqual(ra.evaluate_sustained(seq, last)["status_detail"], "unknown_measurement_prevents_confirmation")

    def test_min_measurements_three_requires_three_consecutive(self):
        T = pd.Timestamp
        conf = {"min_measurements": 3, "min_interval_days": 28, "require_last_below": False}
        seq = [(T("2020-02-01"), 31, True), (T("2020-02-15"), 45, True), (T("2020-03-01"), 60, True)]
        r = ra.evaluate_sustained(seq, conf)
        self.assertEqual((r["status"], r["n_confirming_measurements"], r["confirmation_date"]), ("met_per_protocol", 3, T("2020-03-01")))
        self.assertEqual(ra.evaluate_sustained(seq[:2] + [(T("2020-04-10"), 100, False)], conf)["status_detail"], "transient_not_confirmed_recovered")
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (30, 55, 2000), (45, 55, 2000), (60, 55, 2000)]},
                           {"code": "002", "visits": [(0, 100, 2000), (30, 55, 2000), (60, 55, 2000), (100, 90, 2000)]}])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(lambda p: p["egfr_decline"]["confirmation"].update(min_measurements=3)))
        a = rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((a["status"], a["event_date"], a["confirmation_date"], a["n_confirming_measurements"]), ("met_per_protocol", day(30), day(60), "3"))
        b = rows_for(rows, "002", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((b["status"], b["status_detail"]), ("not_met", "transient_not_confirmed_recovered"))

    def test_excluded_egfr_interrupts_run_instead_of_bridging(self):
        # Review item 3: a recovered value of unverified source (formula tag not recognised)
        # used to be dropped, joining two separate dips into a "sustained" decline.
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, 55, 2000), (120, 95, 2000, "CKD-EPI-2009"), (150, 50, 2000)]},
                           {"code": "002", "visits": [(0, 100, 2000), (90, 55, 2000), (120, 95, 2000), (150, 50, 2000)]}])
        log, rows, qc, out = self.run_protocol(data, self.synth_protocol())
        a = rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((a["status"], a["status_detail"], a["event_date"]), ("not_met", "unknown_measurement_prevents_confirmation", ""))
        self.assertEqual((a["n_followup_measurements"], a["n_unknown_measurements"], a["unknown_measurement_reasons"]), ("2", "1", "egfr_unusable:1"))
        self.assertEqual(a["egfr_source_counts"], "manual:2;excluded:1")
        b = rows_for(rows, "002", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((b["status"], b["status_detail"], b["n_unknown_measurements"]), ("not_met", "met_at_last_measurement_awaiting_confirmation", "0"))
        self.assertIn(("001", "egfr_source_unverified"), {(r["patient_code"], r["code"]) for r in qc})
        flags = {r["id"]: r["egfr_unusable"] for r in read_csv(out / "visits_derived.csv")}
        self.assertEqual((flags["vA001-3"], flags["vA002-3"]), ("True", "False"))

    def test_same_day_measurements_are_averaged(self):
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, 50, 2000), (90, 70, 2000), (150, 60, 2000)]}])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol())
        r = rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]
        # Day 90 is one measurement, mean 60 = exactly 60% of baseline 100.
        self.assertEqual((r["status"], r["event_date"], r["confirmation_date"], r["n_followup_measurements"]), ("met_per_protocol", day(90), day(150), "2"))

    def test_exact_threshold_boundary_is_not_lost_to_floating_point(self):
        # Review item 10: 90*(1-0.30) is 62.99999999999999 in floats, so 63 used to miss.
        self.assertFalse(63 <= 90 * (1 - 30 / 100))
        self.assertTrue(ra.le_pct_of(63, 90, 100 - ra._q(30)))
        data = self.synth([{"code": "001", "visits": [(0, 90, 2000), (90, 63, 2000), (150, 63, 2000)]}])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(lambda p: p["egfr_decline"].update(thresholds_pct_decline=[30])))
        r = rows_for(rows, "001", "egfr_decline_30pct_sustained")[0]
        self.assertEqual((r["status"], r["event_date"]), ("met_per_protocol", day(90)))
        float_mismatch = False
        for pct in (10, 15, 20, 25, 30, 40, 50, 57):
            remaining = 100 - ra._q(pct)
            for base in range(10, 151):
                for v in range(1, base + 1):
                    exact = Fraction(base) * (100 - pct) / 100
                    self.assertEqual(ra.le_pct_of(v, base, remaining), Fraction(v) <= exact, (pct, base, v))
                    self.assertEqual(ra.ge_pct_of(v, base, remaining), Fraction(v) >= exact, (pct, base, v))
                    float_mismatch |= (v <= base * (1 - pct / 100)) != (Fraction(v) <= exact) or (v >= base * (1 - pct / 100)) != (Fraction(v) >= exact)
        self.assertTrue(float_mismatch)  # the float form really was wrong somewhere in this grid
        self.assertTrue(ra.le_pct_of(0.3, 1, 30) and not ra.le_pct_of(0.30000000000000004, 1, 30))

    def test_censoring_date_measurement_is_used(self):
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, 55, 2000), (120, 50, 2000), (200, 90, 2000)],
                            "extra": {"followup_status": "lost", "followup_status_date": day(120)}}],
                          baseline_cols=["followup_status", "followup_status_date"])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(lambda p: p["egfr_decline"]["confirmation"].update(require_last_below=True)))
        r = rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((r["status"], r["event_date"], r["confirmation_date"], r["censor_date"], r["last_measurement_date"]),
                         ("met_per_protocol", day(90), day(120), day(120), day(120)))

    def test_egfr_decline_modules_limits_computation(self):
        path = self.variant_protocol(lambda p: p["egfr_decline"].update(modules=["IGAN"]))
        log, rows, qc, _ = self.run_protocol(DATA, path)
        self.assertEqual([r for r in rows if r["project_id"] == P2], [])
        self.assertTrue(all(r["module"] == "IGAN" for r in rows))
        self.assertIn("eGFR decline not computed for module CKD: excluded by egfr_decline.modules", {r["detail"] for r in qc})
        self.assertEqual(log["protocol"]["term_matches"]["egfr_decline_modules"], {"IGAN": 7})

    # --- baseline rules -----------------------------------------------------------
    def test_baseline_window_mean_and_record_else_window(self):
        def mutate(p):
            p["baseline"] = {"egfr": {"source": "visit_window", "window_days": [-30, 0], "aggregate": "mean"},
                             "upcr": {"source": "baseline_record_else_visit_window", "window_days": [-30, 0], "aggregate": "nearest"}}
        data = self.synth([{"code": "001", "bupcr": "2000", "visits": [(-20, 90, 1500), (0, 110, 1800), (90, 55, 250), (150, 55, 240)]},
                           {"code": "002", "bupcr": "", "visits": [(-20, 90, 1500), (0, 110, 1800), (90, 55, 250), (150, 55, 240)]}])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(mutate))
        e = rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((e["baseline_value"], e["baseline_source"], e["status"], e["n_followup_measurements"]), ("100.0", "visit_window_mean_n2", "met_per_protocol", "2"))
        a = rows_for(rows, "001", "remission_complete_sustained")[0]
        b = rows_for(rows, "002", "remission_complete_sustained")[0]
        self.assertEqual((a["baseline_value"], a["baseline_source"]), ("2000.0", "baseline_record"))
        self.assertEqual((b["baseline_value"], b["baseline_source"], b["status"]), ("1800.0", "visit_window_nearest_day0", "met_per_protocol"))

    # --- remission ------------------------------------------------------------
    def test_igan_complete_and_partial_remission_require_persistence(self):
        cr = self.ep(P1, "A", "001", "remission_complete_sustained")
        self.assertEqual((cr["status"], cr["event_date"], cr["confirmation_date"]), ("met_per_protocol", "2020-03-31", "2020-05-10"))
        self.assertEqual(self.ep(P1, "B", "001", "remission_complete_sustained")["status"], "not_met")
        pr = self.ep(P1, "B", "001", "remission_partial_sustained")
        self.assertEqual((pr["status"], pr["baseline_value"]), ("met_per_protocol", "2000.0"))  # 2.0 g/g converted explicitly
        interrupted = self.ep(P1, "A", "003", "remission_complete_sustained")
        self.assertEqual(interrupted["status"], "not_met")
        self.assertEqual(self.ep(P1, "A", "003", "remission_partial_sustained")["status"], "met_per_protocol")
        self.assertEqual(self.ep(P1, "A", "002", "remission_complete_sustained")["status_detail"], "confirmation_interval_too_short")

    def test_remission_requires_eligible_baseline_upcr(self):
        # Review item 2: baseline UPCR already below the complete threshold, or missing,
        # used to be judged "complete remission".
        data = self.synth([{"code": "001", "bupcr": "200", "visits": [(0, 100, 200), (90, 100, 210), (150, 100, 190)]},
                           {"code": "002", "bupcr": "", "visits": [(0, 100, 250), (90, 100, 210), (150, 100, 190)]},
                           {"code": "003", "bupcr": "800", "visits": [(0, 100, 800), (90, 100, 250), (150, 100, 240)]}])
        _, rows, qc, _ = self.run_protocol(data, self.synth_protocol())
        for level in ("complete", "partial"):
            a = rows_for(rows, "001", f"remission_{level}_sustained")[0]
            self.assertEqual((a["status"], a["status_detail"], a["event_date"]),
                             ("not_evaluable", "baseline_upcr_not_above_remission_threshold:200<=300(complete.upcr_below)", ""))
            b = rows_for(rows, "002", f"remission_{level}_sustained")[0]
            self.assertEqual((b["status"], b["status_detail"]), ("not_evaluable", "baseline_upcr_unavailable:baseline_record_unavailable"))
            self.assertEqual(rows_for(rows, "003", f"remission_{level}_sustained")[0]["status"], "met_per_protocol")
        not_evaluable = {(r["patient_code"], r["detail"]) for r in qc if r["code"] == "endpoint_not_evaluable"}
        self.assertIn(("001", "remission_complete_sustained: baseline_upcr_not_above_remission_threshold:200<=300(complete.upcr_below)"), not_evaluable)
        self.assertIn(("002", "remission_partial_sustained: baseline_upcr_unavailable:baseline_record_unavailable"), not_evaluable)
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(lambda p: p["remission"]["IGAN"].update(baseline_upcr_at_least=1000)))
        c = rows_for(rows, "003", "remission_complete_sustained")[0]
        self.assertEqual((c["status"], c["status_detail"]), ("not_evaluable", "baseline_upcr_below_eligibility:800<1000(baseline_upcr_at_least)"))

    def test_ln_egfr_stability_and_missing_same_day_egfr_is_unknown(self):
        # Review item 4: a UPCR date without same-day eGFR used to count as "not met"
        # (reset); it is now UNKNOWN, like an excluded eGFR: not qualifying, interrupts.
        def ln(p):
            p["remission"]["LN"] = {"requires_pi_confirmation": False, "complete": {"upcr_below": 500, "egfr_stable_max_decline_pct": 15},
                                    "partial": None, "confirmation": {"min_measurements": 2, "min_interval_days": 28, "require_last_below": False},
                                    "assessment_window_days": [90, 365]}
            p["remission"]["IGAN"]["requires_pi_confirmation"] = True
        data = self.synth([{"code": "001", "module": "LN", "visits": [(0, 100, 2000), (100, 95, 300), (130, None, 300), (150, 95, 300)]},
                           {"code": "002", "module": "LN", "visits": [(0, 100, 2000), (100, 95, 300), (130, 95, 300), (150, 95, 300)]},
                           {"code": "003", "module": "LN", "visits": [(0, 100, 2000), (100, 84, 300), (130, 85, 300), (160, 85, 300)]},
                           {"code": "004", "module": "LN", "visits": [(0, 100, 2000), (60, 95, 300), (120, 95, 300), (150, 95, 300)]},
                           {"code": "005", "module": "LN", "visits": [(0, 100, 2000), (100, 95, 300), (130, None, 300, "manual", "200"), (150, 95, 300)]}])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(ln))
        a = rows_for(rows, "001", "remission_complete_sustained")[0]
        self.assertEqual((a["status"], a["status_detail"], a["n_unknown_measurements"], a["unknown_measurement_reasons"]),
                         ("not_met", "unknown_measurement_prevents_confirmation", "1", "no_same_day_egfr:1"))
        b = rows_for(rows, "002", "remission_complete_sustained")[0]
        self.assertEqual((b["status"], b["event_date"], b["confirmation_date"], b["egfr_source_counts"]), ("met_per_protocol", day(100), day(130), "manual:3"))
        # eGFR stability: 84 < 85 (15% below 100) fails; exactly 85 qualifies.
        c = rows_for(rows, "003", "remission_complete_sustained")[0]
        self.assertEqual((c["status"], c["event_date"], c["confirmation_date"]), ("met_per_protocol", day(130), day(160)))
        # Assessment window [90, 365]: the run starting at day 60 counts from its first
        # in-window measurement (day 120), which is the event date.
        d = rows_for(rows, "004", "remission_complete_sustained")[0]
        self.assertEqual((d["status"], d["event_date"], d["confirmation_date"], d["event_days_from_baseline"]), ("met_per_protocol", day(120), day(150), "120"))
        e = rows_for(rows, "005", "remission_complete_sustained")[0]  # creatinine recorded, eGFR not usable
        self.assertEqual((e["status_detail"], e["unknown_measurement_reasons"]), ("unknown_measurement_prevents_confirmation", "same_day_egfr_unusable:1"))

    def test_assessment_window_event_date_is_first_in_window_measurement(self):
        T = pd.Timestamp
        conf = {"min_measurements": 2, "min_interval_days": 28, "require_last_below": False}
        seq = [(T("2020-03-01"), 60, True), (T("2020-04-30"), 120, True), (T("2020-05-30"), 150, True)]
        r = ra.evaluate_sustained(seq, conf, [90, 365])
        self.assertEqual((r["event_date"], r["event_days_from_baseline"], r["confirmation_date"]), (T("2020-04-30"), 120, T("2020-05-30")))
        self.assertEqual(ra.evaluate_sustained(seq, conf, [10, 50])["status_detail"], "met_only_outside_assessment_window")

    def test_upcr_reduction_from_peak_versus_baseline(self):
        # Review item 7: MN partial remission example is a reduction from peak.
        def mn(reference):
            def m(p):
                p["remission"]["MN"] = {"requires_pi_confirmation": False, "complete": {"upcr_below": 300},
                                        "partial": {"upcr_below": 3500, "upcr_reduction_pct_at_least": 50, "upcr_reduction_reference": reference},
                                        "confirmation": {"min_measurements": 2, "min_interval_days": 28, "require_last_below": False}}
                p["remission"]["IGAN"]["requires_pi_confirmation"] = True
            return m
        data = self.synth([{"code": "001", "module": "MN", "bupcr": "4000", "visits": [(0, 100, 4000), (30, 100, 8000), (60, 100, 3400), (90, 100, 3300)]}])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(mn("peak")))
        r = rows_for(rows, "001", "remission_partial_sustained")[0]
        self.assertEqual((r["status"], r["event_date"], r["confirmation_date"]), ("met_per_protocol", day(60), day(90)))
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(mn("baseline")))
        r = rows_for(rows, "001", "remission_partial_sustained")[0]
        self.assertEqual((r["status"], r["status_detail"]), ("not_met", "never_met_criterion"))

    def test_other_criteria_not_evaluated_status(self):
        path = self.variant_protocol(lambda p: p["remission"]["IGAN"].update(additional_criteria_not_computed=["血清白蛋白（合成测试）"]))
        _, rows, _, _ = self.run_protocol(DATA, path)
        r = self.ep(P1, "A", "001", "remission_complete_sustained", rows)
        self.assertEqual((r["status"], r["event_date"], r["criteria_not_evaluated"]),
                         ("upcr_egfr_criteria_met_other_criteria_not_evaluated", "2020-03-31", "血清白蛋白（合成测试）"))
        self.assertEqual(self.ep(P1, "B", "001", "remission_complete_sustained", rows)["status"], "not_met")

    def test_uncovered_and_unconfirmed_modules_are_not_computed_and_reported(self):
        p2 = [r["endpoint"] for r in self.endpoints if r["project_id"] == P2]
        self.assertEqual(sorted(p2), ["egfr_decline_40pct_sustained", "egfr_decline_57pct_sustained"])
        qc = read_csv(self.protocol_out / "qc_issues.csv")
        codes = {(r["code"], r["detail"].split(" ")[0]) for r in qc if r["table"] == "protocol"}
        self.assertIn(("protocol_module_not_covered", "Remission"), codes)
        self.assertIn(("protocol_section_not_confirmed", "remission.MN"), codes)
        self.assertEqual(self.protocol_log["protocol"]["sections_skipped"], {"remission.MN": "requires_pi_confirmation=true"})

    def test_template_alone_computes_nothing_until_pi_confirms(self):
        out = self.tmp / "template-run"
        log = ra.analyze(DATA, out, TEMPLATE / "protocol_template.json")
        self.assertEqual(log["protocol"]["sections_computed"], [])
        self.assertEqual(read_csv(out / "endpoints_per_protocol.csv"), [])
        raw = json.loads((TEMPLATE / "protocol_template.json").read_text(encoding="utf-8"))
        for module in ("IGAN", "LN", "MN"):
            self.assertIs(raw["remission"][module]["requires_pi_confirmation"], True)
            self.assertIn("citation", raw["remission"][module]["reference"])
            self.assertIn("示例", raw["remission"][module]["reference"]["note"])
        self.assertIs(raw["egfr_decline"]["requires_pi_confirmation"], True)

    def test_template_examples_need_pi_and_do_not_assert_guideline_text(self):
        raw = json.loads((TEMPLATE / "protocol_template.json").read_text(encoding="utf-8"))
        ra.validate_protocol(copy.deepcopy(raw))
        self.assertEqual(raw["egfr_decline"]["kidney_failure_or_death_before_confirmation"], "count_as_event")
        self.assertEqual(raw["remission"]["MN"]["partial"]["upcr_reduction_reference"], "peak")
        self.assertEqual(raw["remission"]["IGAN"]["partial"]["upcr_reduction_reference"], "baseline")
        for module in ("IGAN", "LN", "MN"):
            spec = raw["remission"][module]
            self.assertIn("baseline_upcr_at_least", spec)
            self.assertIn("请 PI", spec["reference"]["citation"])
            self.assertNotRegex(spec["reference"]["citation"], r"Chapter|章")
        self.assertIn("常用研究定义示例", raw["remission"]["IGAN"]["reference"]["citation"])
        self.assertIn("KDIGO 2021", raw["remission"]["IGAN"]["reference"]["citation"])
        for module in ("LN", "MN"):
            self.assertIn("baseline_date", raw["remission"][module]["_锚点说明"])
        self.assertIn("治疗开始", raw["_时间锚点说明"])
        readme = (TEMPLATE / "README_PACK.md").read_text(encoding="utf-8")
        self.assertIn("时间锚点的限制", readme)

    def test_confirmed_section_without_named_confirmer_is_refused(self):
        path = self.variant_protocol(lambda p: p["confirmation"].update(confirmed_by=""))
        with self.assertRaises(ra.DataContractError):
            ra.analyze(DATA, self.tmp / "no-confirmer", path)
        self.assertFalse((self.tmp / "no-confirmer").exists())

    def test_misspelled_protocol_key_is_refused(self):
        path = self.variant_protocol(lambda p: p["egfr_decline"]["confirmation"].update(min_interval_day=28))
        with self.assertRaisesRegex(ra.DataContractError, "不认识的字段"):
            ra.load_protocol(path)

    def test_protocol_vocabulary_units_and_required_choices_are_validated(self):
        # Review item 5 (and the new required choices of items 6 and 7).
        def igan_level(level, **kw):
            return lambda p: p["remission"]["IGAN"][level].update(**kw)
        cases = [
            (igan_level("complete", upcr_below=0.3), "合理范围"),  # g/g number in a mg/g protocol
            (lambda p: p["remission"]["IGAN"].update(baseline_upcr_at_least=5), "合理范围"),
            (lambda p: p["remission"]["IGAN"].update(baseline_upcr_at_least=200), "低于完全缓解阈值"),
            (lambda p: p["censoring"].update(event_types=["eskd", "death"]), "不认识的事件类型"),
            (lambda p: p["censoring"].update(event_types=["ESRD"]), "不认识的事件类型"),
            (lambda p: p["remission"].update(IGA=p["remission"].pop("IGAN")), "病种代码只能是"),
            (lambda p: p["egfr_decline"].update(modules=["IgAN"]), "不认识的病种代码"),
            (lambda p: p["remission"]["IGAN"]["partial"].pop("upcr_reduction_reference"), "upcr_reduction_reference"),
            (igan_level("complete", upcr_reduction_reference="peak"), "不应填写"),
            (igan_level("partial", upcr_reduction_reference="nadir"), "upcr_reduction_reference"),
            (lambda p: p["egfr_decline"].pop("kidney_failure_or_death_before_confirmation"), "kidney_failure_or_death_before_confirmation"),
            (lambda p: p["egfr_decline"].update(kidney_failure_or_death_before_confirmation="ignore"), "count_as_event"),
        ]
        for mutate, message in cases:
            with self.subTest(message=message):
                with self.assertRaisesRegex(ra.DataContractError, message):
                    ra.load_protocol(self.variant_protocol(mutate))
        for event_type in ra.EVENT_TYPES:  # the database vocabulary (the events_long CHECK as last set by 0043_biopsy_records.sql) is accepted
            ra.load_protocol(self.variant_protocol(lambda p, t=event_type: p["censoring"].update(event_types=[t])))
        for module in ra.SUPPORTED_MODULES:  # every project module, C3G and ADPKD included, can carry remission and eGFR-decline definitions
            ra.load_protocol(self.variant_protocol(lambda p, m=module: (p["egfr_decline"].update(modules=[m]),
                                                                        p["remission"].update({m: copy.deepcopy(p["remission"]["IGAN"])}))))
        self.assertIn("esrd", ra.EVENT_TYPES)
        self.assertNotIn("eskd", ra.EVENT_TYPES)
        migration = (ROOT / "supabase" / "migrations" / "0043_biopsy_records.sql").read_text(encoding="utf-8")
        check = re.search(r"ADD CONSTRAINT events_long_event_type_check CHECK \(\s*event_type IN \((.*?)\)\);", migration, re.S)
        self.assertEqual(set(re.findall(r"'([a-z0-9_]+)'", check.group(1))), set(ra.EVENT_TYPES), "same vocabulary as the database")
        self.assertNotIn("vaccination", ra.EVENT_TYPES, "0043 withdrew the vaccination event type")
        # The project creation guard as last re-created (0046_adpkd_module.sql) holds the module whitelist.
        modules_sql = (ROOT / "supabase" / "migrations" / "0046_adpkd_module.sql").read_text(encoding="utf-8")
        self.assertIn("if new.module not in (" + ",".join(f"'{m}'" for m in ra.SUPPORTED_MODULES) + ")", modules_sql)

    def test_confirmed_terms_without_matching_records_are_refused_unless_allowed(self):
        cases = [
            (lambda p: p["remission"]["MN"].update(requires_pi_confirmation=False), "remission.MN"),
            (lambda p: p["egfr_decline"].update(modules=["IGAN", "DKD"]), "egfr_decline.modules DKD"),
            (lambda p: p["censoring"].update(event_types=["esrd", "death", "custom"]), "censoring.event_types custom"),
        ]
        for mutate, message in cases:
            with self.subTest(message=message):
                path = self.variant_protocol(mutate)
                out = self.newdir() / "out"
                with self.assertRaisesRegex(ra.DataContractError, message + ".*--allow-unmatched-protocol-terms"):
                    ra.analyze(DATA, out, path)
                self.assertFalse(out.exists())
        log = ra.analyze(DATA, self.newdir() / "out", self.variant_protocol(cases[2][0]), allow_unmatched_protocol_terms=True)
        tm = log["protocol"]["term_matches"]
        self.assertEqual(tm["event_types"], {"custom": {"records": 0, "confirmed": 0}, "death": {"records": 2, "confirmed": 1},
                                             "esrd": {"records": 1, "confirmed": 1}})
        self.assertEqual((tm["unmatched"], tm["proceeded_with_allow_unmatched_protocol_terms"]), (["censoring.event_types custom（事件表无此类型记录）"], True))
        self.assertEqual(log["overrides"]["allow_unmatched_protocol_terms"], True)
        main = self.protocol_log["protocol"]["term_matches"]
        self.assertEqual((main["remission_modules"], main["unmatched"]), ({"IGAN": 7}, []))

    # --- kidney failure / death before confirmation -----------------------------
    def test_kidney_failure_or_death_before_confirmation_count_as_event(self):
        # Review item 6: censoring at esrd/death is informative and under-counts decline.
        path = self.variant_protocol(lambda p: p["egfr_decline"].update(kidney_failure_or_death_before_confirmation="count_as_event"))
        log, rows, _, _ = self.run_protocol(DATA, path)
        name = "egfr_decline_40pct_sustained_or_kidney_failure_death"
        died = self.ep(P1, "A", "007", name, rows)
        self.assertEqual((died["status"], died["status_detail"], died["event_basis"], died["event_date"], died["event_days_from_baseline"]),
                         ("met_per_protocol", "kidney_failure_or_death_before_confirmation", "event_death", "2020-05-20", "140"))
        decline = self.ep(P2, "A", "001", name, rows)  # esrd at d200 after the decline was confirmed at d150
        self.assertEqual((decline["event_basis"], decline["event_date"], decline["confirmation_date"]), ("egfr_measurements", "2020-03-31", "2020-05-30"))
        esrd = self.ep(P2, "A", "001", "egfr_decline_57pct_sustained_or_kidney_failure_death", rows)
        self.assertEqual((esrd["status"], esrd["event_basis"], esrd["event_date"]), ("met_per_protocol", "event_esrd", "2020-07-19"))
        lost = self.ep(P1, "B", "002", name, rows)  # loss to follow-up is not a kidney failure/death event
        self.assertEqual((lost["status"], lost["censor_reason"]), ("not_met", "followup_status_lost"))
        self.assertEqual(log["protocol"]["egfr_decline_kidney_failure_or_death_before_confirmation"], "count_as_event")
        self.assertFalse([r for r in rows if r["endpoint"] == "egfr_decline_40pct_sustained"])
        # censor (the fixture's choice): the same patients are censored, not events.
        censored = self.ep(P1, "A", "007", "egfr_decline_40pct_sustained")
        self.assertEqual((censored["status"], censored["censor_date"], censored["event_basis"]), ("not_met", "2020-05-20", ""))
        p2 = self.ep(P2, "A", "001", "egfr_decline_57pct_sustained")
        self.assertEqual((p2["status"], p2["censor_reason"]), ("not_met", "event_esrd"))

    def test_count_as_event_uses_confirmed_events_only_and_needs_baseline(self):
        def mutate(p):
            p["egfr_decline"]["kidney_failure_or_death_before_confirmation"] = "count_as_event"
            p["censoring"]["event_types"] = ["esrd"]
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, 55, 2000), (150, 50, 2000)],
                            "extra": {"followup_status": "deceased", "followup_status_date": day(120)}},
                           {"code": "002", "visits": [(10, 100, 2000), (90, 55, 2000)], "extra": {"followup_status": "active"}},
                           {"code": "003", "visits": [(0, 100, 2000), (90, 55, 2000)], "extra": {"followup_status": "active"}}],
                          events=[{"code": "002", "type": "esrd", "day": 100}, {"code": "003", "type": "esrd", "day": 100, "confirmed": "false"}],
                          baseline_cols=["followup_status", "followup_status_date"])
        _, rows, _, _ = self.run_protocol(data, self.synth_protocol(mutate))
        name = "egfr_decline_40pct_sustained_or_kidney_failure_death"
        status_only = rows_for(rows, "001", name)[0]  # follow-up status is administrative: censors, not an event
        self.assertEqual((status_only["status"], status_only["censor_reason"], status_only["event_basis"]), ("not_met", "followup_status_deceased", ""))
        no_baseline = rows_for(rows, "002", name)[0]  # an event does not make an unevaluable patient count
        self.assertEqual((no_baseline["status"], no_baseline["status_detail"]), ("not_evaluable", "baseline_egfr_unavailable:no_measurement_in_baseline_window"))
        unconfirmed = rows_for(rows, "003", name)[0]
        self.assertEqual((unconfirmed["status"], unconfirmed["censor_date"]), ("not_met", ""))

    # --- censoring ------------------------------------------------------------
    def test_followup_status_and_confirmed_events_censor_later_measurements(self):
        lost = self.ep(P1, "B", "002", "egfr_decline_40pct_sustained")
        self.assertEqual((lost["status"], lost["censor_date"], lost["censor_reason"]), ("not_met", "2020-04-30", "followup_status_lost"))
        died = self.ep(P1, "A", "007", "egfr_decline_40pct_sustained")
        self.assertEqual((died["status"], died["censor_reason"]), ("not_met", "event_death"))
        # Unconfirmed death event for A/001 does not censor.
        self.assertEqual(self.ep(P1, "A", "001", "remission_complete_sustained")["censor_date"], "")

    def test_followup_status_date_before_baseline_is_ignored_and_reported(self):
        # Review item 8: such a date used to censor every measurement including baseline.
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, 50, 2000), (150, 50, 2000)],
                            "extra": {"followup_status": "lost", "followup_status_date": "2019-06-01"}}],
                          events=[{"code": "001", "type": "death", "day": -10}], baseline_cols=["followup_status", "followup_status_date"])
        _, rows, qc, _ = self.run_protocol(data, self.synth_protocol(lambda p: p["censoring"].update(event_types=["death"])))
        r = rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]
        self.assertEqual((r["status"], r["event_date"], r["censor_date"], r["censor_reason"]),
                         ("met_per_protocol", day(90), "", "followup_status_lost_date_before_baseline_ignored;event_death_not_after_baseline_ignored"))
        codes = {r["code"] for r in qc if r["patient_code"] == "001"}
        self.assertLessEqual({"followup_status_date_before_baseline", "censoring_event_not_after_baseline"}, codes)

    def test_not_evaluable_endpoints_are_written_to_qc(self):
        data = self.synth([{"code": "001", "visits": [(5, 100, 2000), (90, 50, 2000), (150, 50, 2000)]},
                           {"code": "002", "baseline_date": "2020/01/01", "visits": [(90, 50, 2000)]}])
        _, rows, qc, _ = self.run_protocol(data, self.synth_protocol())
        self.assertEqual(rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]["status_detail"], "baseline_egfr_unavailable:no_measurement_in_baseline_window")
        details = {(r["patient_code"], r["detail"]) for r in qc if r["code"] == "endpoint_not_evaluable"}
        self.assertIn(("001", "egfr_decline_40pct_sustained: baseline_egfr_unavailable:no_measurement_in_baseline_window"), details)
        self.assertIn(("002", "remission_complete_sustained: baseline_date_unavailable"), details)
        self.assertIn(("002", "missing_or_invalid_baseline_date"), {(r["patient_code"], r["code"]) for r in qc})

    def test_missing_followup_status_columns_are_noted(self):
        data = self.variant_data(mutate_baseline=lambda h, rows: ([c for c in h if not c.startswith("followup_status")], rows))
        log = ra.analyze(data, self.tmp / "no-followup-status", PROTOCOL)
        self.assertEqual(log["followup_status"]["columns_present"], [])
        self.assertIn("未提供随访状态", log["followup_status"]["note"])
        rows = read_csv(self.tmp / "no-followup-status" / "endpoints_per_protocol.csv")
        lost = [r for r in rows if (r["center_code"], r["patient_code"], r["endpoint"]) == ("B", "002", "egfr_decline_40pct_sustained")][0]
        self.assertEqual(lost["status"], "met_per_protocol")  # without status it is not censored

    # --- timepoints -------------------------------------------------------------
    def test_timepoints_select_nearest_visit_within_window_after_censoring(self):
        rows = read_csv(self.protocol_out / "timepoints_per_protocol.csv")
        a = [r for r in rows if (r["project_id"], r["center_code"], r["patient_code"]) == (P1, "A", "001")]
        self.assertEqual(len(a), 1)
        self.assertEqual({k: a[0][k] for k in ["nominal_day", "window_days", "status", "visit_id", "visit_date", "days_from_baseline", "egfr", "upcr",
                                               "egfr_baseline_protocol", "upcr_baseline_protocol", "egfr_pct_change", "upcr_pct_change", "tie_count"]},
                         {"nominal_day": "180", "window_days": "60", "status": "visit_selected", "visit_id": "v001-4", "visit_date": "2020-07-19",
                          "days_from_baseline": "200.0", "egfr": "92.0", "upcr": "270.0", "egfr_baseline_protocol": "100.0",
                          "upcr_baseline_protocol": "2000.0", "egfr_pct_change": "-8.0", "upcr_pct_change": "-86.5", "tie_count": "1"})
        lost = [r for r in rows if (r["project_id"], r["center_code"], r["patient_code"]) == (P1, "B", "002")][0]
        self.assertEqual((lost["status"], lost["visit_id"]), ("no_visit_in_window", ""))  # d150/d200 visits are after censoring at d120

    # --- units ----------------------------------------------------------------
    def test_protocol_unit_differing_from_data_unit_is_refused(self):
        def to_gg(p):  # a self-consistent g/g protocol: valid on its own, but not the data unit
            p["units"]["upcr"] = "g/g"
            p["remission"]["IGAN"]["complete"]["upcr_below"] = 0.3
            p["remission"]["IGAN"]["partial"]["upcr_below"] = 1.0
            p["remission"]["MN"]["complete"]["upcr_below"] = 0.3
        path = self.variant_protocol(to_gg)
        with self.assertRaisesRegex(ra.DataContractError, "不一致"):
            ra.analyze(DATA, self.tmp / "unit-protocol", path)
        self.assertFalse((self.tmp / "unit-protocol").exists())
        # Declaring g/g while keeping mg/g numbers is refused at validation (magnitude).
        with self.assertRaisesRegex(ra.DataContractError, "合理范围"):
            ra.load_protocol(self.variant_protocol(lambda p: p["units"].update(upcr="g/g")))

    def test_explicit_data_unit_differing_from_protocol_is_refused(self):
        def add_unit(header, rows):
            for r in rows:
                r["upcr_unit"] = "g/g"
            return header + ["upcr_unit"], rows
        with self.assertRaisesRegex(ra.DataContractError, "visits_long.upcr_unit"):
            ra.analyze(self.variant_data(mutate_visits=add_unit), self.tmp / "unit-column", PROTOCOL)
        with self.assertRaisesRegex(ra.DataContractError, "data_dictionary.json"):
            ra.analyze(self.variant_data(dictionary={"visits": {"upcr": "g/g"}}), self.tmp / "unit-dictionary", PROTOCOL)
        ok = ra.analyze(self.variant_data(dictionary={"visits": {"upcr": "mg/g"}}), self.tmp / "unit-dictionary-ok", PROTOCOL)
        self.assertEqual(ok["protocol"]["visit_upcr_unit_declarations"], [{"unit": "mg/g", "source": "data_dictionary.json"}])

    # --- without protocol -----------------------------------------------------
    def test_without_protocol_only_unconfirmed_candidates_are_written(self):
        files = set(self.candidate_log["generated_files"])
        self.assertIn("candidates_unconfirmed_egfr_decline.csv", files)
        self.assertIn("candidates_unconfirmed_igan_proteinuria.csv", files)
        self.assertFalse({"endpoints_per_protocol.csv", "candidates_egfr_decline.csv", "candidates_igan_proteinuria.csv"} & files)
        for name in ["candidates_unconfirmed_egfr_decline.csv", "candidates_unconfirmed_igan_proteinuria.csv"]:
            with (self.candidate_out / name).open(encoding="utf-8-sig", newline="") as f:
                header = next(csv.reader(f))
            self.assertTrue(all(c in ra.KEY or c.startswith("candidate_unconfirmed_") for c in header), header)
        statuses = {r["candidate_unconfirmed_egfr_le_60pct_of_baseline_status"] for r in read_csv(self.candidate_out / "candidates_unconfirmed_egfr_decline.csv")}
        self.assertLessEqual(statuses, {"candidate_unconfirmed", "not_observed_in_available_followup", "not_evaluable"})
        self.assertEqual((self.candidate_log["mode"], self.candidate_log["protocol"]), ("candidates_only", None))
        model = json.loads((self.candidate_out / "MODEL_STATUS.json").read_text(encoding="utf-8"))
        self.assertEqual(model["clinical_endpoint_adjudication"]["status"], "not_executed_unconfirmed_candidates_only")

    def test_table1_has_no_endpoint_rates(self):
        from openpyxl import load_workbook
        for out in (self.candidate_out, self.protocol_out):
            ws = load_workbook(out / "table1_baseline.xlsx").active
            labels = [str(row[0]).lower() for row in ws.iter_rows(min_row=2, values_only=True)]
            for banned in ("remission", "decline", "缓解", "下降", "candidate", "endpoint"):
                self.assertFalse([x for x in labels if banned in x], (out.name, banned))

    # --- identity and consent ---------------------------------------------------
    def test_same_patient_code_across_centers_and_projects_is_not_merged(self):
        code001 = {(r["project_id"], r["center_code"]) for r in self.endpoints if r["patient_code"] == "001"}
        self.assertEqual(code001, {(P1, "A"), (P1, "B"), (P2, "A")})
        a = self.ep(P1, "A", "001", "egfr_decline_40pct_sustained")
        b = self.ep(P1, "B", "001", "egfr_decline_40pct_sustained")
        self.assertNotEqual(a["status"], b["status"])
        self.assertEqual(a["n_followup_measurements"], "3")
        self.assertEqual(b["n_followup_measurements"], "3")

    def test_consent_filter_counts_and_excludes_all_records(self):
        c = self.protocol_log["consent_filter"]
        self.assertEqual((c["column_present"], c["patients_before_n"], c["included_n"], c["excluded_n"]), (True, 10, 8, 2))
        self.assertEqual(c["excluded_by_status"], {"declined": 1, "unconfirmed": 1})
        for out in (self.protocol_out, self.candidate_out):
            for name in ["patients_derived.csv", "visits_derived.csv"]:
                codes = {(r["center_code"], r["patient_code"]) for r in read_csv(out / name)}
                self.assertFalse({("A", "005"), ("A", "006")} & codes)
        self.assertEqual(self.candidate_log["consent_filter"]["excluded_n"], 2)
        self.assertEqual(c["by_center"], [
            {"project_id": P1, "center_code": "A", "patients_n": 7, "included_n": 5, "excluded_n": 2,
             "excluded_by_status": {"declined": 1, "unconfirmed": 1}, "consent_status_column_absent_in_source_export": False},
            {"project_id": P1, "center_code": "B", "patients_n": 2, "included_n": 2, "excluded_n": 0, "excluded_by_status": {},
             "consent_status_column_absent_in_source_export": False},
            {"project_id": P2, "center_code": "A", "patients_n": 1, "included_n": 1, "excluded_n": 0, "excluded_by_status": {},
             "consent_status_column_absent_in_source_export": False}])
        qc = read_csv(self.protocol_out / "qc_issues.csv")
        self.assertIn(("A", "", "consent_filter_center_counts", 'included 5/7; excluded 2 by consent_status {"declined": 1, "unconfirmed": 1}'),
                      {(r["center_code"], r["patient_code"], r["code"], r["detail"]) for r in qc if r["project_id"] == P1})

    def test_consent_filter_removes_excluded_patients_from_every_output_table(self):
        extra = {
            "labs_long.csv": (ra.TABLE_COLUMNS["labs_long"], [
                {"project_id": P1, "center_code": "A", "patient_code": code, "id": f"l{code}", "module": "IGAN", "lab_date": "2020-02-01",
                 "lab_test_code": "ALB", "lab_name": "Albumin", "lab_value": "4.0", "lab_unit": "g/dL"} for code in ("001", "005", "006")]),
            "meds_long.csv": (ra.TABLE_COLUMNS["meds_long"], [
                {"project_id": P1, "center_code": "A", "patient_code": code, "id": f"m{code}", "module": "IGAN", "drug_name": "x",
                 "drug_class": "RASi", "start_date": "2020-01-01"} for code in ("001", "005")]),
            "events_long.csv": (ra.TABLE_COLUMNS["events_long"], [
                {"project_id": P1, "center_code": "A", "patient_code": "005", "id": "e005", "module": "IGAN", "event_type": "death",
                 "event_date": "2020-03-01", "confirmed": "true", "source": "staff"}]),
        }
        data = self.variant_data(extra_tables=extra)
        excluded = {(P1, "A", "005"), (P1, "A", "006")}
        for protocol in (PROTOCOL, None):
            out = self.newdir() / "out"
            ra.analyze(data, out, protocol)
            checked = 0
            for path in sorted(out.glob("*.csv")):
                rows = read_csv(path)
                if rows and set(ra.KEY) <= set(rows[0]):
                    checked += 1
                    self.assertFalse(excluded & {tuple(r[k] for k in ra.KEY) for r in rows}, path.name)
            self.assertIn((P1, "A", "001"), {tuple(r[k] for k in ra.KEY) for r in read_csv(out / "labs_review.csv")})
            self.assertEqual(checked, 9)  # every keyed CSV output was checked

    def test_without_consent_column_nothing_is_filtered_and_log_says_so(self):
        data = self.variant_data(mutate_baseline=lambda h, rows: ([c for c in h if not c.startswith("consent_")], rows))
        log = ra.analyze(data, self.tmp / "no-consent")
        self.assertEqual((log["consent_filter"]["column_present"], log["consent_filter"]["excluded_n"], log["patients_n"]), (False, 0, 10))
        self.assertIn("未提供同意状态", log["consent_filter"]["note"])

    def test_blank_or_default_consent_for_everyone_stops_the_run(self):
        for value, label in (("", "(empty)"), ("unconfirmed", "unconfirmed")):
            with self.subTest(value=value):
                data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, 50, 200)], "extra": {"consent_status": value}}],
                                  baseline_cols=["consent_status"])
                out = self.newdir() / "out"
                with self.assertRaises(ra.DataContractError) as ctx:
                    ra.analyze(data, out, self.synth_protocol())
                message = str(ctx.exception)
                for needle in ("纳入 0", label, "数据库默认值 unconfirmed", "在工作台", "--allow-fully-excluded-centers", "English"):
                    self.assertIn(needle, message)
                self.assertFalse(out.exists())
                # With nobody included the confirmed IGAN definition also matches no patient,
                # so proceeding needs both explicit overrides.
                with self.assertRaisesRegex(ra.DataContractError, "remission.IGAN"):
                    ra.analyze(data, out, self.synth_protocol(), allow_fully_excluded_centers=True)
                log = ra.analyze(data, out, self.synth_protocol(), allow_fully_excluded_centers=True, allow_unmatched_protocol_terms=True)
                self.assertEqual((log["patients_n"], log["consent_filter"]["included_n"], log["consent_filter"]["excluded_by_status"]), (0, 0, {label: 1}))
                self.assertTrue(log["consent_filter"]["proceeded_with_allow_fully_excluded_centers"])
                self.assertEqual(read_csv(out / "endpoints_per_protocol.csv"), [])

    def test_analysis_cli_reports_failure_with_nonzero_exit(self):
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000)], "extra": {"consent_status": "unconfirmed"}}], baseline_cols=["consent_status"])
        stderr = io.StringIO()
        with mock.patch.object(sys, "argv", ["run_analysis.py", "--data", str(data), "--out", str(self.newdir() / "out")]), \
                contextlib.redirect_stderr(stderr), self.assertRaises(SystemExit) as ctx:
            ra.main()
        self.assertEqual(ctx.exception.code, 2)
        self.assertIn("ANALYSIS_FAILED", stderr.getvalue())
        self.assertIn("unconfirmed", stderr.getvalue())

    # --- multi-center merge -------------------------------------------------------
    def centers_with_missing_columns(self):
        root = self.newdir()
        status_cols = ["consent_status", "followup_status", "followup_status_date"]
        a = make_data(root / "A", [{"code": c, "center": "A", "visits": [(0, 100, 2000), (90, 50, 200), (150, 50, 200)],
                                    "extra": {"consent_status": "consented", "followup_status": "active"}} for c in ("001", "002")],
                      baseline_cols=status_cols)
        b = make_data(root / "B", [{"code": c, "center": "B", "visits": [(0, 100, 2000), (90, 50, 200), (150, 50, 200)]} for c in ("001", "002", "003")])
        return root, a, b

    def test_merge_refuses_centers_with_missing_columns(self):
        # Review item 1: the center without consent_status used to be blank-filled and
        # then excluded entirely while the run still reported success.
        root, a, b = self.centers_with_missing_columns()
        out, qc = root / "merged", root / "merge_qc.xlsx"
        with self.assertRaises(ra.DataContractError) as ctx:
            mc.merge_all([a, b], out, qc, verbose=False)
        message = str(ctx.exception)
        self.assertIn(f"{b.resolve()} 的 patients_baseline.csv 缺少列：consent_status, followup_status, followup_status_date", message)
        self.assertIn("--allow-missing-columns", message)
        self.assertNotIn(str(a.resolve()) + " 的", message)
        self.assertFalse(out.exists())
        from openpyxl import load_workbook
        sheet = list(load_workbook(qc)["ColumnMismatch"].iter_rows(values_only=True))
        self.assertEqual(sheet[1][:3], (str(b.resolve()), "patients_baseline", "consent_status, followup_status, followup_status_date"))
        stderr = io.StringIO()
        with contextlib.redirect_stderr(stderr), self.assertRaises(SystemExit) as cli:
            mc.main(["--dirs", str(a), str(b), "--out", str(root / "merged-cli"), "--qc", str(root / "qc-cli.xlsx")])
        self.assertEqual(cli.exception.code, 2)
        self.assertIn("MERGE_FAILED", stderr.getvalue())
        self.assertIn("consent_status", stderr.getvalue())

    def test_merge_with_allowed_missing_columns_is_recorded_and_analysis_reports_per_center(self):
        root, a, b = self.centers_with_missing_columns()
        merged = root / "merged"
        mc.merge_all([a, b], merged, root / "merge_qc.xlsx", verbose=False, allow_missing_columns=True)
        manifest = json.loads((merged / "MERGE_MANIFEST.json").read_text(encoding="utf-8"))
        self.assertTrue(manifest["allow_missing_columns"])
        self.assertEqual(manifest["column_mismatch"], [{"source_dir": str(b.resolve()), "table": "patients_baseline",
                                                        "missing_columns": ["consent_status", "followup_status", "followup_status_date"],
                                                        "critical_missing": ["consent_status", "followup_status", "followup_status_date"]}])
        protocol = self.synth_protocol()
        with self.assertRaises(ra.DataContractError) as ctx:
            ra.analyze(merged, root / "out-refused", protocol)
        message = str(ctx.exception)
        self.assertIn(f"项目 {SP} / 中心 B：3 名患者全部被排除（(empty) 3）；该中心的源导出缺少 consent_status 列", message)
        self.assertNotIn("中心 A：", message)
        log = ra.analyze(merged, root / "out", protocol, allow_fully_excluded_centers=True)
        by_center = {c["center_code"]: c for c in log["consent_filter"]["by_center"]}
        self.assertEqual((by_center["A"]["included_n"], by_center["B"]["included_n"], by_center["B"]["excluded_by_status"],
                          by_center["B"]["consent_status_column_absent_in_source_export"]), (2, 0, {"(empty)": 3}, True))
        self.assertEqual(log["consent_filter"]["fully_excluded_centers"], [f"{SP}/B"])
        self.assertEqual(log["followup_status"]["centers_without_followup_status_in_source_export"],
                         [{"project_id": SP, "center_code": "B", "missing_columns": ["followup_status", "followup_status_date"]}])
        self.assertIn(f"{SP}/B", log["followup_status"]["note"])
        self.assertTrue(log["merge_manifest"]["allow_missing_columns"])
        self.assertEqual(log["merge_manifest"]["sha256"], hashlib.sha256((merged / "MERGE_MANIFEST.json").read_bytes()).hexdigest())
        qc = read_csv(root / "out" / "qc_issues.csv")
        self.assertIn(("B", "consent_filter_center_fully_excluded"), {(r["center_code"], r["code"]) for r in qc})
        self.assertEqual({r["center_code"] for r in read_csv(root / "out" / "endpoints_per_protocol.csv")}, {"A"})
        methods = (root / "out" / "METHODS_ACTUAL_EN.md").read_text(encoding="utf-8")
        self.assertIn(f"{SP}/B", methods)

    def test_merge_of_identical_column_sets_still_succeeds(self):
        root = self.newdir()
        a = make_data(root / "A", [{"code": "001", "center": "A", "visits": [(0, 100, 2000)]}])
        b = make_data(root / "B", [{"code": "001", "center": "B", "visits": [(0, 100, 2000)]}])
        mc.merge_all([a, b], root / "merged", root / "qc.xlsx", verbose=False)
        manifest = json.loads((root / "merged" / "MERGE_MANIFEST.json").read_text(encoding="utf-8"))
        self.assertEqual((manifest["status"], manifest["column_mismatch"], manifest["allow_missing_columns"]), ("merged_no_conflicts", [], False))

    # --- provenance -----------------------------------------------------------
    def test_protocol_text_and_hash_recorded_in_run_log_and_methods(self):
        digest = hashlib.sha256(PROTOCOL.read_bytes()).hexdigest()
        text = PROTOCOL.read_text(encoding="utf-8")
        log = json.loads((self.protocol_out / "RUN_LOG.json").read_text(encoding="utf-8"))
        self.assertEqual((log["protocol"]["sha256"], log["protocol"]["text"], log["mode"]), (digest, text, "protocol"))
        methods = (self.protocol_out / "METHODS_ACTUAL_EN.md").read_text(encoding="utf-8")
        self.assertIn(digest, methods)
        self.assertIn(text.rstrip(), methods)
        self.assertIn("TEST-SYNTHETIC@1.0", methods)
        self.assertEqual((self.protocol_out / "protocol_used.json").read_bytes(), PROTOCOL.read_bytes())

    def test_protocol_used_is_the_validated_bytes(self):
        original = ra.load_protocol
        captured = {}

        def spy(path):
            result = original(path)
            captured["bytes"] = result[4]
            return result
        with mock.patch.object(ra, "load_protocol", spy):
            out = self.newdir() / "out"
            ra.analyze(DATA, out, PROTOCOL)
        self.assertEqual((out / "protocol_used.json").read_bytes(), captured["bytes"])
        self.assertEqual(hashlib.sha256(captured["bytes"]).hexdigest(), json.loads((out / "RUN_LOG.json").read_text(encoding="utf-8"))["protocol"]["sha256"])

    def test_run_log_records_package_provenance_environment_and_module_counts(self):
        pkg = self.newdir() / "pkg"
        data = pkg / "analysis" / "data"
        shutil.copytree(DATA, data)
        (pkg / "data_dictionary.json").write_text(json.dumps({"visits": {"upcr": "mg/g"}}), encoding="utf-8")
        (pkg / "research_data.json").write_text("{}", encoding="utf-8")
        sha = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in data.glob("*.csv")}
        files = {f"analysis/data/{name}": {"sha256": digest} for name, digest in sha.items()}
        files["analysis/data/events_long.csv"] = {"sha256": "0" * 64}
        (pkg / "export_manifest.json").write_text(json.dumps({"format": "registry-export-v3", "project_id": P1, "center_code": "A",
                                                             "exported_at": "2026-01-01T00:00:00Z", "metadata_register_id": "SNAP-1",
                                                             "source_content_sha256": "a" * 64, "files": files}), encoding="utf-8")
        log = ra.analyze(data, self.newdir() / "out", PROTOCOL)
        self.assertEqual(log["data_dictionary_sha256"], hashlib.sha256((pkg / "data_dictionary.json").read_bytes()).hexdigest())
        manifest = log["provenance_files"]["export_manifest.json"]
        self.assertEqual((manifest["location"], manifest["sha256"]), ("../../export_manifest.json", hashlib.sha256((pkg / "export_manifest.json").read_bytes()).hexdigest()))
        self.assertEqual((manifest["identifiers"]["metadata_register_id"], manifest["identifiers"]["project_id"]), ("SNAP-1", P1))
        self.assertEqual(manifest["input_csv_check"]["visits_long.csv"], "matches_export_manifest")
        self.assertEqual(manifest["input_csv_check"]["events_long.csv"], "differs_from_export_manifest")
        self.assertEqual(log["provenance_files"]["research_data.json"]["sha256"], hashlib.sha256(b"{}").hexdigest())
        self.assertEqual(log["environment"]["python"], platform.python_version())
        self.assertEqual(set(log["dependencies"]), {"pandas", "numpy", "openpyxl", "matplotlib"})
        counts = self.protocol_log["protocol"]["endpoint_status_counts_by_module"]
        self.assertEqual(set(counts), {"IGAN", "CKD"})
        self.assertEqual(counts["CKD"], {"egfr_decline_40pct_sustained|met_per_protocol": 1, "egfr_decline_57pct_sustained|not_met": 1})
        self.assertEqual(self.protocol_log["provenance_files"]["export_manifest.json"], None)

    def test_egfr_source_composition_is_reported(self):
        # Review item 11: which eGFRs were recomputed, manual, or excluded.
        stored = f"{ra.ckdepi2021(90 / 88.4, 50, 'M'):.2f}"
        data = self.synth([{"code": "001", "visits": [(0, 100, 2000), (90, stored, 2000, "CKD-EPI-2021-Cr", "90"), (100, None, 2000, "CKD-EPI-2021-Cr", "90"),
                                                      (120, 95, 2000, "CKD-EPI-2009"), (150, 50, 2000), (160, None, 2000, "manual", "")]}])
        log, rows, _, _ = self.run_protocol(data, self.synth_protocol())
        self.assertEqual(rows_for(rows, "001", "egfr_decline_40pct_sustained")[0]["egfr_source_counts"], "recomputed:2;manual:1;excluded:1")
        sources = log["egfr_sources"]
        self.assertEqual({k: sources[k] for k in sources if k.endswith("_n")},
                         {"dated_visits_n": 6, "recomputed_from_creatinine_n": 2, "recomputed_stored_value_agrees_n": 1,
                          "recomputed_stored_value_differs_n": 0, "recomputed_no_stored_value_n": 1, "recomputed_schwartz_2009_n": 0, "manual_value_used_n": 2,
                          "excluded_recorded_but_unusable_n": 1, "no_egfr_or_creatinine_recorded_n": 1})


if __name__ == "__main__":
    unittest.main()
