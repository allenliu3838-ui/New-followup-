"""KidneySphere reproducible analysis (integrity-v3-protocol).

No database/network access. Strict project/record identity; original CSVs are never
changed.

Without --protocol: descriptive review only; fixed numerical thresholds are emitted as
UNCONFIRMED candidates (file names and headers say so), never as endpoints or rates.
With --protocol <json>: sustained endpoints are determined only from the parameters in
that study protocol file (PI-confirmed sections only); its full text and SHA-256 are
recorded in RUN_LOG.json and METHODS_ACTUAL_EN.md. This is algorithmic determination
per protocol, not manual clinical adjudication. Recorded-but-excluded measurements are
UNKNOWN and interrupt confirmation runs; threshold comparisons are exact.

Population: when patients_baseline has consent_status, only consented/waiver patients
are analysed (the registry default 'unconfirmed' is excluded). The run stops if any
project/center, or the whole input, has no included patient, unless
--allow-fully-excluded-centers is given. Confirmed protocol modules/event types with
no matching record stop the run unless --allow-unmatched-protocol-terms is given.
KFRE and cohort inferential models remain disabled pending protocol validation.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import math
import os
import platform
import re
from fractions import Fraction
from pathlib import Path
import shutil
import sys
import tempfile
from datetime import datetime, timezone

import numpy as np
import pandas as pd

VERSION = "integrity-v3-protocol"
BASE = Path(__file__).resolve().parent
KEY = ["project_id", "center_code", "patient_code"]
TABLE_COLUMNS = {
    "patients_baseline": KEY + ["id", "module", "sex", "birth_year", "baseline_date", "baseline_scr", "baseline_upcr", "baseline_upcr_unit"],
    "visits_long": KEY + ["id", "module", "visit_date", "sbp", "dbp", "scr_umol_l", "upcr", "egfr", "egfr_formula_version", "height_cm"],
    "labs_long": KEY + ["id", "module", "lab_date", "lab_test_code", "lab_name", "lab_value", "lab_unit", "value_raw", "unit_symbol", "value_standard", "standard_unit", "measured_at"],
    "meds_long": KEY + ["id", "module", "drug_name", "drug_class", "dose", "route", "frequency", "start_date", "end_date"],
    "variants_long": KEY + ["id", "module", "test_date", "gene", "variant", "hgvs_c", "hgvs_p"],
    "events_long": KEY + ["id", "module", "event_type", "event_date", "confirmed", "source"],
}
EGFR_REFERENCE = "https://escholarship.org/content/qt3gj2d8m2/qt3gj2d8m2.pdf"


class DataContractError(ValueError):
    """Input is ambiguous; do not emit an apparently successful analysis."""


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def number(value):
    try:
        result = float(value)
        return result if math.isfinite(result) else np.nan
    except (TypeError, ValueError):
        return np.nan


def _q(x):
    """Exact rational of a finite number read at its shortest decimal form.

    63 stays 63 and 0.3 stays 3/10, so a threshold boundary is not lost to binary
    floating point (in floats 90*(1-0.30) is 62.99999999999999, and 63 would miss a
    30% decline). Fractions pass through unchanged.
    """
    if isinstance(x, Fraction):
        return x
    if isinstance(x, int) and not isinstance(x, bool):
        return Fraction(x)
    return Fraction(repr(float(x)))


def le_pct_of(value, reference, pct):
    """value <= reference * pct / 100, compared exactly (no float rounding)."""
    return bool(_q(value) * 100 <= _q(reference) * _q(pct))


def ge_pct_of(value, reference, pct):
    """value >= reference * pct / 100, compared exactly (no float rounding)."""
    return bool(_q(value) * 100 >= _q(reference) * _q(pct))


def date(value):
    if value is None or pd.isna(value) or not str(value).strip():
        return pd.NaT
    # Registry dates are ISO calendar dates. Do not guess locale/day-month order.
    if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", str(value)):
        return pd.NaT
    try:
        return pd.Timestamp(datetime.strptime(str(value), "%Y-%m-%d"))
    except ValueError:
        return pd.NaT


def ckdepi2021(scr_mg_dl, age, sex):
    """Adult creatinine equation: Inker 2021 Table 2; age here is year-derived.

    Missing/unknown sex, nonpositive creatinine, or age outside 18..120 => NaN.
    This does not validate suitability for an individual clinical application.
    """
    scr, age = number(scr_mg_dl), number(age)
    sx = str(sex).strip().upper()
    if not (np.isfinite(scr) and scr > 0 and np.isfinite(age) and 18 <= age <= 120 and sx in {"M", "F"}):
        return np.nan
    k, alpha = (0.7, -0.241) if sx == "F" else (0.9, -0.302)
    return 142 * min(scr/k, 1)**alpha * max(scr/k, 1)**-1.200 * 0.9938**age * (1.012 if sx == "F" else 1)


def schwartz2009(scr_mg_dl, height_cm, age):
    """Bedside Schwartz 2009 (0.413 x height cm / creatinine mg/dL) for children 1-17
    years by year-derived age; the registry stores it as 'Schwartz-2009-bedside'.

    Missing/invalid height (30-250 cm), nonpositive creatinine or an age outside 1..17
    => NaN. Not validated for infants or for adults.
    """
    scr, height, age = number(scr_mg_dl), number(height_cm), number(age)
    if not (np.isfinite(scr) and scr > 0 and np.isfinite(height) and 30 <= height <= 250 and np.isfinite(age) and 1 <= age < 18):
        return np.nan
    return 0.413 * height / scr


def issue(issues, table, row, code, detail):
    issues.append({"table": table, **{k: row.get(k, "") for k in KEY}, "record_id": row.get("id", ""), "code": code, "detail": detail})


def read_table(path, table, required=False, allow_duplicates=False):
    """No numeric inference for identifiers; even 'NA' is a literal identifier.

    df.attrs["source_columns"] keeps the CSV header as exported (None when the file is
    absent), before any default column is added, so callers can see which columns a
    source export really had.
    """
    path = Path(path)
    if not path.exists():
        if required:
            raise DataContractError(f"Missing required {path.name}")
        df = pd.DataFrame(columns=TABLE_COLUMNS[table])
        df.attrs["source_columns"] = None
        return df
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fields = reader.fieldnames
        if not fields or len(fields) != len(set(fields)):
            raise DataContractError(f"{path.name}: missing or duplicate CSV header")
        rows = list(reader)
    if any(None in r or any(v is None for v in r.values()) for r in rows):
        raise DataContractError(f"{path.name}: malformed CSV row width")
    df = pd.DataFrame(rows, columns=fields, dtype=object)
    if not df.empty:
        missing = set(KEY + ["id", "module"]) - set(df.columns)
        if missing:
            raise DataContractError(f"{path.name}: missing identity columns {sorted(missing)}; re-export using the updated registry. Never infer project identity.")
        for col in KEY + ["id", "module"]:
            if df[col].map(lambda v: not v or v != v.strip()).any():
                raise DataContractError(f"{path.name}: blank or surrounding whitespace in {col}; correct source explicitly")
        if not allow_duplicates and df.duplicated(["project_id", "id"]).any():
            raise DataContractError(f"{path.name}: duplicate record IDs; resolve or merge exact retransmissions first")
        if not allow_duplicates and table == "patients_baseline" and df.duplicated(KEY).any():
            raise DataContractError("patients_baseline.csv: multiple baselines for one project/center/patient")
    for col in TABLE_COLUMNS[table]:
        if col not in df:
            df[col] = ""
    df.attrs["source_columns"] = list(fields)
    return df


def load_tables(data_dir):
    frames = {t: read_table(Path(data_dir)/f"{t}.csv", t, t in {"patients_baseline", "visits_long"}) for t in TABLE_COLUMNS}
    baseline = frames["patients_baseline"]
    identity = {tuple(r[k] for k in KEY): r["module"] for r in baseline.to_dict("records")}
    # A project/center must be mapped consistently to the same study module.
    if not baseline.empty and (baseline.groupby(["project_id", "center_code"])["module"].nunique() > 1).any():
        raise DataContractError("One project/center has conflicting disease modules")
    for table, df in frames.items():
        if table == "patients_baseline":
            continue
        for r in df.to_dict("records"):
            key = tuple(r[k] for k in KEY)
            if key not in identity:
                raise DataContractError(f"{table}: orphan record {r['id']}; no matching project/center/patient baseline")
            if r["module"] != identity[key]:
                raise DataContractError(f"{table}: module differs from patient baseline for record {r['id']}")
    return frames


def upcr_mg_g(value, unit):
    n = number(value)
    converted = n * {"mg/g": 1, "g/g": 1000}.get(str(unit), np.nan) if np.isfinite(n) and n >= 0 else np.nan
    return converted if np.isfinite(converted) else np.nan


def derive(frames, issues):
    patients = frames["patients_baseline"].copy()
    for c in ["baseline_anchor", "age_baseline", "baseline_egfr", "baseline_upcr_final"]:
        patients[c] = pd.Series(index=patients.index, dtype=object)
    for idx, r in patients.iterrows():
        anchor = date(r["baseline_date"])
        by = number(r["birth_year"])
        age = anchor.year - by if pd.notna(anchor) and np.isfinite(by) and by.is_integer() else np.nan
        if np.isfinite(age) and not 0 <= age <= 120:
            issue(issues, "patients_baseline", r, "invalid_baseline_age", "Impossible year-derived age excluded from age summaries")
            age = np.nan
        patients.at[idx, "baseline_anchor"] = anchor
        patients.at[idx, "age_baseline"] = age
        patients.at[idx, "baseline_egfr"] = ckdepi2021(number(r["baseline_scr"])/88.4, age, r["sex"])
        if not np.isfinite(patients.at[idx, "baseline_egfr"]):
            issue(issues, "patients_baseline", r, "baseline_egfr_not_evaluable", "Requires positive creatinine, valid baseline date, M/F sex, and adult year-derived age; no later-visit substitution")
        patients.at[idx, "baseline_upcr_final"] = upcr_mg_g(r["baseline_upcr"], r["baseline_upcr_unit"])
        if pd.isna(anchor):
            issue(issues, "patients_baseline", r, "missing_or_invalid_baseline_date", "No first-visit substitution; anchored outcomes not evaluable")
        if r["baseline_upcr"] and not np.isfinite(patients.at[idx, "baseline_upcr_final"]):
            issue(issues, "patients_baseline", r, "unverified_baseline_upcr", "UPCR excluded from change/ratio calculations until its original unit and value are verified")
        if str(r["sex"]).upper() not in {"M", "F"}:
            issue(issues, "patients_baseline", r, "sex_missing_or_unknown", "No sex-specific eGFR calculation; unknown is not male")
    v = frames["visits_long"].copy()
    v["egfr_recorded"] = v["egfr"]
    v["visit_date_raw"] = v["visit_date"]
    v["visit_date"] = v["visit_date"].map(date)
    recorded_text = {col: v[col].map(lambda x: str(x).strip() != "") for col in ["egfr", "scr_umol_l", "upcr"]}
    v["_kidney_function_recorded"] = recorded_text["egfr"] | recorded_text["scr_umol_l"]
    for col in ["sbp", "dbp", "scr_umol_l", "upcr", "height_cm"]:
        raw = v[col].copy()
        v[col] = raw.map(number)
        for idx in raw.index:
            if raw[idx] and not np.isfinite(v.at[idx, col]):
                issue(issues, "visits_long", v.loc[idx], "invalid_numeric", col)
            if np.isfinite(v.at[idx, col]) and v.at[idx, col] < 0:
                issue(issues, "visits_long", v.loc[idx], "negative_measurement", col)
                v.at[idx, col] = np.nan
    # A recorded but unusable UPCR is an unknown measurement, not an absent one.
    v["upcr_unusable"] = recorded_text["upcr"] & v["upcr"].isna()
    v = v.merge(patients[KEY + ["baseline_anchor", "baseline_egfr", "baseline_upcr_final", "sex", "birth_year"]], on=KEY, how="left", validate="many_to_one")
    v["baseline_anchor"] = pd.to_datetime(v["baseline_anchor"])
    v["visit_date"] = pd.to_datetime(v["visit_date"])
    v["days_from_baseline"] = (v["visit_date"]-v["baseline_anchor"]).dt.days
    v["time_yr"] = v["days_from_baseline"] / 365.25
    v["egfr"] = np.nan
    v["egfr_analysis_source"] = "not_evaluable"
    for idx, r in v.iterrows():
        if pd.notna(r["days_from_baseline"]) and r["days_from_baseline"] < 0:
            issue(issues, "visits_long", r, "visit_before_baseline", "Excluded from post-baseline outcomes, slopes and trends")
        if pd.isna(r["visit_date"]):
            issue(issues, "visits_long", r, "missing_or_invalid_visit_date", "Excluded from dated outcomes")
        recorded = number(r["egfr_recorded"])
        version = r["egfr_formula_version"]
        if version == "manual" and np.isfinite(recorded) and recorded > 0:
            v.at[idx, "egfr"] = recorded
            v.at[idx, "egfr_analysis_source"] = "recorded_manual"
        elif version in {"CKD-EPI-2021-Cr", "Schwartz-2009-bedside"}:
            by = number(r["birth_year"])
            age = r["visit_date"].year-by if pd.notna(r["visit_date"]) and np.isfinite(by) and by.is_integer() else np.nan
            if version == "CKD-EPI-2021-Cr":
                calculated = ckdepi2021(r["scr_umol_l"]/88.4, age, r["sex"])
                unavailable = "Current creatinine/date/sex/adult age cannot support recalculation"
            else:
                # Paediatric tag: recomputed from the visit's own height, like the registry trigger.
                calculated = schwartz2009(r["scr_umol_l"]/88.4, r["height_cm"], age)
                unavailable = "Current creatinine/date/visit height/paediatric age (1-17) cannot support Schwartz recalculation"
            v.at[idx, "egfr"] = calculated
            if not np.isfinite(calculated):
                issue(issues, "visits_long", r, "egfr_inputs_unavailable", unavailable)
            if np.isfinite(calculated):
                v.at[idx, "egfr_analysis_source"] = RECOMPUTED_EGFR if version == "CKD-EPI-2021-Cr" else RECOMPUTED_EGFR_PAEDIATRIC
                if np.isfinite(recorded) and abs(calculated-recorded) > 0.11:
                    issue(issues, "visits_long", r, "stored_egfr_mismatch", "Analysis recomputed from current creatinine/demographics; original retained in egfr_recorded")
        elif np.isfinite(recorded):
            issue(issues, "visits_long", r, "egfr_source_unverified", "Value retained as egfr_recorded, excluded from derived analysis until source is verified")
    # An eGFR or creatinine was recorded for the visit but no usable analysis eGFR
    # results (unverified source, invalid value, missing demographics): in protocol
    # endpoint sequences this visit is an UNKNOWN measurement that interrupts runs; it
    # is never silently skipped (which would bridge two separate dips into a "run").
    v["egfr_unusable"] = v["_kidney_function_recorded"].astype(bool) & v["egfr"].isna()
    v = v.drop(columns=["_kidney_function_recorded"])
    return patients, v


# Explicit per-row conversions; no inference from magnitude or cohort distribution.
LAB_UNITS = {
    "ALB": ("g/dL", {"g/dL": 1, "g/L": 0.1}),
    "CREAT": ("mg/dL", {"mg/dL": 1, "μmol/L": 1/88.4, "umol/L": 1/88.4}),
    "UPCR": ("mg/g", {"mg/g": 1, "g/g": 1000}),
    "UACR": ("mg/g", {"mg/g": 1}),
    "PHOS": ("mg/dL", {"mg/dL": 1, "mmol/L": 3.097}),
    "CA": ("mg/dL", {"mg/dL": 1, "mmol/L": 4.008}),
    "CO2": ("mmol/L", {"mmol/L": 1, "mEq/L": 1}),
    "CD19": ("cells/μL", {"cells/μL": 1}),
    "BKV": ("copies/mL", {"copies/mL": 1}),
    "CMV": ("IU/mL", {"IU/mL": 1}),
}


def normalize_labs(df, issues):
    out = df.copy()
    out["analysis_value"] = np.nan
    out["analysis_unit"] = ""
    out["analysis_status"] = "unverified"
    for idx, r in out.iterrows():
        code = r["lab_test_code"]
        has_raw = r["value_raw"] != ""
        has_unit = r["unit_symbol"] != ""
        if has_raw != has_unit:
            issue(issues, "labs_long", r, "incomplete_raw_unit_pair", "Partial structured value/unit; never mix with a legacy value/unit pair")
            continue
        if has_raw:
            raw, unit = number(r["value_raw"]), r["unit_symbol"]
        else:
            raw, unit = number(r["lab_value"]), r["lab_unit"]
        std, su = number(r["value_standard"]), r["standard_unit"]
        if (r["value_standard"] != "" and not np.isfinite(std)) or ((r["value_standard"] != "") != (su != "")):
            issue(issues, "labs_long", r, "invalid_standard_value_unit_pair", "Stored standard value/unit is malformed or partial; excluded pending correction")
            continue
        target, factors = LAB_UNITS.get(code, (su, {su: 1} if su else {}))
        if not np.isfinite(raw) or raw < 0 or unit not in factors:
            issue(issues, "labs_long", r, "lab_conversion_unverified", "Original value/unit retained; no guessed standard value")
            continue
        expected = raw*factors[unit]
        if not np.isfinite(expected):
            issue(issues, "labs_long", r, "converted_value_nonfinite", "Unit conversion overflow; no usable analysis value")
            continue
        if np.isfinite(std):
            if su not in factors or not math.isclose(std*factors[su], expected, rel_tol=0.001, abs_tol=0.001):
                issue(issues, "labs_long", r, "lab_standard_mismatch", "Stored standard value/unit conflicts with original; excluded pending correction")
                continue
        out.at[idx, "analysis_value"] = expected
        out.at[idx, "analysis_unit"] = target
        out.at[idx, "analysis_status"] = "explicit_unit_conversion"
    return out


def choose_12month(patients, visits):
    rows = []
    for r in patients.to_dict("records"):
        sub = visits
        for k in KEY:
            sub = sub[sub[k] == r[k]]
        sub = sub[sub.days_from_baseline.between(270, 450)].copy()
        result = {k: r[k] for k in KEY}
        result.update(baseline_record_id=r["id"], baseline_date=r["baseline_anchor"], egfr_baseline=r["baseline_egfr"], upcr_baseline=r["baseline_upcr_final"], status="no_evaluable_annual_visit", visit_id="", visit_12m_date=pd.NaT, days_from_baseline=np.nan, egfr_12m=np.nan, upcr_12m=np.nan, scr_12m_umol_l=np.nan, egfr_delta=np.nan, upcr_delta=np.nan, tie_count=0)
        if pd.isna(r["baseline_anchor"]):
            result["status"] = "baseline_date_unavailable"
        elif not sub.empty:
            sub["distance"] = (sub.days_from_baseline-365).abs()
            sub = sub.sort_values(["distance", "visit_date", "id"], kind="stable")
            picked = sub.iloc[0]  # One entire row; missing values stay missing.
            result.update(status="visit_selected", visit_id=picked["id"], visit_12m_date=picked.visit_date, days_from_baseline=picked.days_from_baseline, egfr_12m=picked.egfr, upcr_12m=picked.upcr, scr_12m_umol_l=picked.scr_umol_l, egfr_delta=picked.egfr-number(r["baseline_egfr"]), upcr_delta=picked.upcr-number(r["baseline_upcr_final"]), tie_count=int((sub.distance == picked.distance).sum()))
        rows.append(result)
    return pd.DataFrame(rows, columns=KEY + ["baseline_record_id", "baseline_date", "egfr_baseline", "upcr_baseline", "status", "visit_id", "visit_12m_date", "days_from_baseline", "egfr_12m", "upcr_12m", "scr_12m_umol_l", "egfr_delta", "upcr_delta", "tie_count"])


CANDIDATE_EGFR_LABELS = [(60, "egfr_le_60pct_of_baseline"), (43, "egfr_le_43pct_of_baseline")]
CANDIDATE_UPCR_LABELS = ["upcr_below_300", "upcr_half_and_below_1000"]


def candidates(patients, visits):
    """UNCONFIRMED single-measurement threshold candidates (used only without --protocol).

    Fixed descriptive thresholds; no persistence/confirmation; not endpoints, not
    remission, never summarised as rates. Every output column says candidate_unconfirmed.
    """
    declines, protein = [], []
    for r in patients.to_dict("records"):
        sub = visits
        for k in KEY:
            sub = sub[sub[k] == r[k]]
        sub = sub[sub.days_from_baseline > 0].sort_values(["visit_date", "id"], kind="stable")
        dr = {k:r[k] for k in KEY}
        pr = dict(dr)
        for pct_of_baseline, label in CANDIDATE_EGFR_LABELS:
            base = number(r["baseline_egfr"])
            valid = sub[sub.egfr.notna()] if np.isfinite(base) and base > 0 else sub.iloc[0:0]
            hit = valid[valid.egfr.map(lambda x: le_pct_of(x, base, pct_of_baseline))] if not valid.empty else valid
            col = f"candidate_unconfirmed_{label}"
            dr[col] = bool(not hit.empty) if not valid.empty else None
            dr[col+"_status"] = "candidate_unconfirmed" if not hit.empty else ("not_observed_in_available_followup" if not valid.empty else "not_evaluable")
            dr[col+"_visit_id"] = hit.iloc[0]["id"] if not hit.empty else ""
            dr[col+"_date"] = hit.iloc[0].visit_date if not hit.empty else pd.NaT
        declines.append(dr)
        if str(r["module"]).upper() != "IGAN":
            continue
        valid = sub[sub.upcr.notna()]
        base = number(r["baseline_upcr_final"])
        for label in CANDIDATE_UPCR_LABELS:
            eligible = valid if label == "upcr_below_300" or (np.isfinite(base) and base > 0) else valid.iloc[0:0]
            if eligible.empty:
                hit = eligible
            elif label == "upcr_below_300":
                hit = eligible[eligible.upcr < 300]
            else:
                hit = eligible[eligible.upcr.map(lambda x: le_pct_of(x, base, 50)) & (eligible.upcr < 1000)]
            col = f"candidate_unconfirmed_{label}"
            pr[col] = bool(not hit.empty) if not eligible.empty else None
            pr[col+"_status"] = "candidate_unconfirmed" if not hit.empty else ("not_observed_in_available_followup" if not eligible.empty else "not_evaluable")
            pr[col+"_visit_id"] = hit.iloc[0]["id"] if not hit.empty else ""
        protein.append(pr)
    dcols=KEY+[f"candidate_unconfirmed_{label}{s}" for _,label in CANDIDATE_EGFR_LABELS for s in ["","_status","_visit_id","_date"]]
    pcols=KEY+[f"candidate_unconfirmed_{label}{s}" for label in CANDIDATE_UPCR_LABELS for s in ["","_status","_visit_id"]]
    return pd.DataFrame(declines,columns=dcols), pd.DataFrame(protein,columns=pcols)


# ---------------------------------------------------------------------------
# Study-protocol parameterised endpoints (--protocol). Nothing here has a default
# clinical threshold: every number comes from the protocol file, and a section is
# only computed after the protocol states the PI confirmed it.
# ---------------------------------------------------------------------------
PROTOCOL_SCHEMA = "kidneysphere-analysis-protocol/1"
PROTOCOL_MAX_BYTES = 1_000_000
# Units the analysis compares against. Visit UPCR is stored in mg/g by the registry
# export contract (data_dictionary.json: visits.upcr); baseline UPCR is converted per
# row only from its explicit baseline_upcr_unit. Thresholds are never rescaled.
VISIT_UPCR_UNIT = "mg/g"
UNIT_ALIASES = {
    "upcr": {"mg/g": "mg/g", "g/g": "g/g", "mg/mmol": "mg/mmol"},
    "egfr": {"mL/min/1.73m2": "mL/min/1.73m2", "mL/min/1.73m²": "mL/min/1.73m2"},
    "scr": {"umol/L": "umol/L", "μmol/L": "umol/L", "µmol/L": "umol/L", "mg/dL": "mg/dL"},
}
ANALYSIS_UNITS = {"upcr": VISIT_UPCR_UNIT, "egfr": "mL/min/1.73m2", "scr": "umol/L"}
# Plausible magnitude of a UPCR threshold in each declared unit. A value outside the
# range is almost always a unit slip (0.3 written for 0.3 g/g in a mg/g protocol).
UPCR_THRESHOLD_RANGES = {"mg/g": (10, 100000), "g/g": (0.01, 100), "mg/mmol": (1, 11300)}
CONSENT_INCLUDED = ("consented", "waiver")
CONSENT_VALUES = ("unconfirmed", "consented", "waiver", "declined", "withdrawn")
FOLLOWUP_STATUS_VALUES = ("active", "lost", "withdrawn", "deceased", "kidney_failure", "transferred")
# Registry vocabularies (supabase/migrations): module codes accepted for projects
# (0033_billing_registration.sql) and the events_long.event_type CHECK (0002_events.sql).
SUPPORTED_MODULES = ("IGAN", "LN", "MN", "DKD", "CKD", "KTX", "GENERAL")
EVENT_TYPES = ("egfr_decline_40pct", "egfr_decline_57pct", "esrd", "death", "complete_remission", "partial_remission", "custom",
               "dialysis_start", "kidney_transplant", "relapse")
# Kidney failure (esrd, start of dialysis, kidney transplant) or death.
KF_DEATH_EVENT_TYPES = ("esrd", "dialysis_start", "kidney_transplant", "death")
KF_DEATH_OPTIONS = ("count_as_event", "censor")
UPCR_REDUCTION_REFERENCES = ("baseline", "peak")
RECOMPUTED_EGFR = "recomputed_CKD-EPI-2021-Cr_year_age"
RECOMPUTED_EGFR_PAEDIATRIC = "recomputed_Schwartz-2009-bedside_year_age"
RECOMPUTED_EGFR_SOURCES = (RECOMPUTED_EGFR, RECOMPUTED_EGFR_PAEDIATRIC)
ENDPOINT_COLUMNS = KEY + ["module", "protocol", "determination", "endpoint", "status", "status_detail", "event_basis",
                          "event_date", "confirmation_date", "event_days_from_baseline", "n_confirming_measurements",
                          "baseline_value", "baseline_unit", "baseline_source", "n_followup_measurements",
                          "n_unknown_measurements", "unknown_measurement_reasons", "egfr_source_counts",
                          "last_measurement_date", "last_measurement_days_from_baseline", "censor_date",
                          "censor_reason", "criteria_not_evaluated"]
TIMEPOINT_COLUMNS = KEY + ["module", "protocol", "nominal_day", "window_days", "status", "visit_id", "visit_date",
                           "days_from_baseline", "egfr", "upcr", "egfr_baseline_protocol", "upcr_baseline_protocol",
                           "egfr_pct_change", "upcr_pct_change", "tie_count"]
CONSENT_HINT_ZH = "数据库默认值 unconfirmed（未确认）与空值都会被排除，需先在工作台登记知情同意状态（consented 已同意 或 waiver 伦理批准豁免）后重新导出。"
CONSENT_HINT_EN = ("The registry default consent_status 'unconfirmed' and blank values are excluded; record consent "
                   "(consented or waiver) in the registry workbench and re-export.")


def _pfail(where, message):
    raise DataContractError(f"研究方案文件 {where}：{message}")


def _check_keys(obj, where, required, optional=()):
    if not isinstance(obj, dict):
        _pfail(where, "应为 JSON 对象")
    keys = {k for k in obj if not str(k).startswith("_")}  # keys beginning "_" are comments
    missing = [k for k in required if k not in obj]
    unknown = sorted(keys - set(required) - set(optional))
    if missing:
        _pfail(where, f"缺少字段 {missing}")
    if unknown:
        _pfail(where, f"不认识的字段 {unknown}（拼写错误会让规则静默失效，因此拒绝运行）")


def _pnum(value, where, lo=None, hi=None, integer=False, allow_none=False, lo_open=False, hi_open=False):
    if value is None and allow_none:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        _pfail(where, "应为有限数值")
    if integer and not float(value).is_integer():
        _pfail(where, "应为整数")
    if lo is not None and (value <= lo if lo_open else value < lo):
        _pfail(where, f"超出允许范围（下限 {lo}）")
    if hi is not None and (value >= hi if hi_open else value > hi):
        _pfail(where, f"超出允许范围（上限 {hi}）")
    return int(value) if integer else float(value)


def _pbool(value, where):
    if not isinstance(value, bool):
        _pfail(where, "应为 true 或 false")
    return value


def _ptext(value, where, max_len=200, allow_empty=False):
    if not isinstance(value, str) or (not allow_empty and not value.strip()) or len(value) > max_len:
        _pfail(where, f"应为非空文字（≤{max_len} 字符）")
    return value


def _pwindow(value, where, allow_none=False):
    if value is None and allow_none:
        return None
    if not isinstance(value, list) or len(value) != 2:
        _pfail(where, "应为 [起始天, 结束天]")
    a = _pnum(value[0], where, -3650, 36500, integer=True)
    b = _pnum(value[1], where, -3650, 36500, integer=True)
    if a > b:
        _pfail(where, "起始天不能大于结束天")
    return [a, b]


def _pupcr(value, where, unit, allow_none=False):
    """A UPCR threshold in the protocol's declared unit, within a plausible magnitude."""
    if value is None and allow_none:
        return None
    v = _pnum(value, where)
    lo, hi = UPCR_THRESHOLD_RANGES[unit]
    if not lo <= v <= hi:
        _pfail(where, f"UPCR 数值 {v:g} 不在 units.upcr={unit} 的合理范围 [{lo:g}, {hi:g}] 内，数值量级与声明的单位不符"
                      f"（0.3 g/g = 300 mg/g ≈ 34 mg/mmol；mg/g 方案里写 0.3 通常是把 g/g 数值填错了单位）。"
                      "脚本不会猜测或换算单位，请 PI 核对后修改")
    return v


def _validate_confirmation(conf, where):
    _check_keys(conf, where, ["min_measurements", "min_interval_days", "require_last_below"])
    out = {"min_measurements": _pnum(conf["min_measurements"], where+".min_measurements", 1, 20, integer=True),
           "min_interval_days": _pnum(conf["min_interval_days"], where+".min_interval_days", 0, 3650, integer=True),
           "require_last_below": _pbool(conf["require_last_below"], where+".require_last_below")}
    if out["min_measurements"] == 1 and out["min_interval_days"] > 0:
        _pfail(where, "min_measurements=1 时不存在第二次测量，min_interval_days 必须为 0")
    return out


def _validate_baseline_rule(rule, where):
    _check_keys(rule, where, ["source"], ["window_days", "aggregate"])
    source = rule["source"]
    if source not in ("baseline_record", "visit_window", "baseline_record_else_visit_window"):
        _pfail(where+".source", "只能是 baseline_record / visit_window / baseline_record_else_visit_window")
    out = {"source": source, "window_days": None, "aggregate": None}
    if source != "baseline_record":
        out["window_days"] = _pwindow(rule.get("window_days"), where+".window_days")
        if out["window_days"][0] < -365 or out["window_days"][1] > 365:
            _pfail(where+".window_days", "基线窗口应在基线日期前后 365 天内")
        if rule.get("aggregate") not in ("nearest", "mean"):
            _pfail(where+".aggregate", "只能是 nearest 或 mean")
        out["aggregate"] = rule["aggregate"]
    elif "window_days" in rule or "aggregate" in rule:
        _pfail(where, "source=baseline_record 时不使用 window_days/aggregate，请删除以免误解")
    return out


def _validate_reference(ref, where):
    _check_keys(ref, where, ["citation"], ["note", "url"])
    _ptext(ref["citation"], where+".citation", 1000)
    return ref


def _validate_level(level, where, upcr_unit):
    _check_keys(level, where, [], ["upcr_below", "upcr_reduction_pct_at_least", "upcr_reduction_reference", "egfr_stable_max_decline_pct"])
    out = {"upcr_below": _pupcr(level.get("upcr_below"), where+".upcr_below", upcr_unit, allow_none=True),
           "upcr_reduction_pct_at_least": _pnum(level.get("upcr_reduction_pct_at_least"), where+".upcr_reduction_pct_at_least", 0, 100, allow_none=True, lo_open=True, hi_open=True),
           "egfr_stable_max_decline_pct": _pnum(level.get("egfr_stable_max_decline_pct"), where+".egfr_stable_max_decline_pct", 0, 100, allow_none=True, hi_open=True)}
    if out["upcr_below"] is None and out["upcr_reduction_pct_at_least"] is None:
        _pfail(where, "至少需要 upcr_below 或 upcr_reduction_pct_at_least 之一")
    ref = level.get("upcr_reduction_reference")
    if out["upcr_reduction_pct_at_least"] is not None:
        if ref not in UPCR_REDUCTION_REFERENCES:
            _pfail(where+".upcr_reduction_reference", "设置 upcr_reduction_pct_at_least 时必须写明下降的参照：baseline（方案基线值）"
                                                      "或 peak（基线窗口起至该次测量之前的最高 UPCR，含基线值）")
    elif ref is not None:
        _pfail(where+".upcr_reduction_reference", "未设置 upcr_reduction_pct_at_least 时不应填写，请删除以免误解")
    out["upcr_reduction_reference"] = ref
    return out


def validate_protocol(raw):
    """Strictly validate a parsed protocol; returns a normalised copy used by the run."""
    _check_keys(raw, "顶层", ["schema", "protocol_id", "version", "confirmation", "units", "baseline", "visit_window_days",
                             "timepoints_days", "egfr_decline", "remission", "censoring"], ["title", "description", "status"])
    if raw["schema"] != PROTOCOL_SCHEMA:
        _pfail("schema", f"应为 {PROTOCOL_SCHEMA}")
    pid = _ptext(raw["protocol_id"], "protocol_id", 80)
    ver = _ptext(raw["version"], "version", 40)
    for name, value in [("protocol_id", pid), ("version", ver)]:
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", value):
            _pfail(name, "只能包含字母、数字、点、下划线、连字符")
    p = {"protocol_id": pid, "version": ver}
    conf = raw["confirmation"]
    _check_keys(conf, "confirmation", ["confirmed_by", "confirmed_date"], ["role", "statement"])
    p["confirmation"] = {"confirmed_by": _ptext(conf["confirmed_by"], "confirmation.confirmed_by", 120, allow_empty=True),
                         "confirmed_date": _ptext(conf["confirmed_date"], "confirmation.confirmed_date", 10, allow_empty=True),
                         "role": conf.get("role", ""), "statement": conf.get("statement", "")}
    units = raw["units"]
    _check_keys(units, "units", ["upcr", "egfr"], ["scr"])
    p["units"] = {}
    for key, value in units.items():
        if str(key).startswith("_"):
            continue
        canonical = UNIT_ALIASES[key].get(value)
        if canonical is None:
            _pfail(f"units.{key}", f"不支持的单位 {value!r}")
        p["units"][key] = canonical
    upcr_unit = p["units"]["upcr"]
    base = raw["baseline"]
    _check_keys(base, "baseline", ["egfr", "upcr"])
    p["baseline"] = {m: _validate_baseline_rule(base[m], f"baseline.{m}") for m in ("egfr", "upcr")}
    p["visit_window_days"] = _pnum(raw["visit_window_days"], "visit_window_days", 0, 365, integer=True)
    if not isinstance(raw["timepoints_days"], list):
        _pfail("timepoints_days", "应为天数列表（可为空列表）")
    p["timepoints_days"] = sorted({_pnum(x, "timepoints_days", 1, 36500, integer=True) for x in raw["timepoints_days"]})
    ed = raw["egfr_decline"]
    if ed is None:
        p["egfr_decline"] = None
    else:
        _check_keys(ed, "egfr_decline", ["requires_pi_confirmation", "thresholds_pct_decline", "confirmation",
                                         "kidney_failure_or_death_before_confirmation"], ["reference", "modules"])
        if not isinstance(ed["thresholds_pct_decline"], list) or not ed["thresholds_pct_decline"]:
            _pfail("egfr_decline.thresholds_pct_decline", "应为非空列表，如 [40]")
        thresholds = sorted({_pnum(x, "egfr_decline.thresholds_pct_decline", 0, 100, lo_open=True, hi_open=True) for x in ed["thresholds_pct_decline"]})
        modules = ed.get("modules")
        if modules is not None and (not isinstance(modules, list) or not modules or not all(isinstance(m, str) for m in modules)):
            _pfail("egfr_decline.modules", "应为病种代码列表或省略（省略表示全部）")
        bad = sorted({m for m in modules or [] if m not in SUPPORTED_MODULES})
        if bad:
            _pfail("egfr_decline.modules", f"不认识的病种代码 {bad}；只能是登记系统的 module 取值 {list(SUPPORTED_MODULES)}（区分大小写）")
        kf = ed["kidney_failure_or_death_before_confirmation"]
        if kf not in KF_DEATH_OPTIONS:
            _pfail("egfr_decline.kidney_failure_or_death_before_confirmation",
                   "只能是 count_as_event（确认下降之前发生已确认的 esrd/dialysis_start/kidney_transplant/death 事件时计为终点事件，即复合终点）或 censor（在该日期截尾）")
        p["egfr_decline"] = {"requires_pi_confirmation": _pbool(ed["requires_pi_confirmation"], "egfr_decline.requires_pi_confirmation"),
                             "thresholds_pct_decline": thresholds,
                             "confirmation": _validate_confirmation(ed["confirmation"], "egfr_decline.confirmation"),
                             "kidney_failure_or_death_before_confirmation": kf,
                             "modules": None if modules is None else sorted(set(modules)),
                             "reference": _validate_reference(ed["reference"], "egfr_decline.reference") if "reference" in ed else None}
    rem = raw["remission"]
    if not isinstance(rem, dict):
        _pfail("remission", "应为按病种代码组织的对象（可为空对象）")
    p["remission"] = {}
    for module, spec in rem.items():
        if str(module).startswith("_"):
            continue
        where = f"remission.{module}"
        if module not in SUPPORTED_MODULES:
            _pfail(where, f"病种代码只能是登记系统的 module 取值 {list(SUPPORTED_MODULES)}（区分大小写，例如 IgA 肾病写 IGAN）")
        _check_keys(spec, where, ["requires_pi_confirmation", "confirmation", "complete", "partial"],
                    ["reference", "assessment_window_days", "additional_criteria_not_computed", "label", "baseline_upcr_at_least"])
        extra = spec.get("additional_criteria_not_computed", [])
        if not isinstance(extra, list) or not all(isinstance(x, str) and x.strip() for x in extra):
            _pfail(where+".additional_criteria_not_computed", "应为文字列表")
        levels = {}
        for level in ("complete", "partial"):
            levels[level] = None if spec[level] is None else _validate_level(spec[level], f"{where}.{level}", upcr_unit)
        if levels["complete"] is None and levels["partial"] is None:
            _pfail(where, "complete 与 partial 不能都为 null")
        at_least = _pupcr(spec.get("baseline_upcr_at_least"), where+".baseline_upcr_at_least", upcr_unit, allow_none=True)
        complete_cut = levels["complete"]["upcr_below"] if levels["complete"] else None
        if at_least is not None and complete_cut is not None and at_least < complete_cut:
            _pfail(where+".baseline_upcr_at_least", f"{at_least:g} 低于完全缓解阈值 complete.upcr_below={complete_cut:g}："
                                                    "基线已处于缓解范围的患者会被判为缓解；请 PI 核对")
        p["remission"][module] = {"requires_pi_confirmation": _pbool(spec["requires_pi_confirmation"], where+".requires_pi_confirmation"),
                                  "confirmation": _validate_confirmation(spec["confirmation"], where+".confirmation"),
                                  "assessment_window_days": _pwindow(spec.get("assessment_window_days"), where+".assessment_window_days", allow_none=True),
                                  "additional_criteria_not_computed": list(extra), **levels, "baseline_upcr_at_least": at_least,
                                  "reference": _validate_reference(spec["reference"], where+".reference") if "reference" in spec else None}
    cen = raw["censoring"]
    _check_keys(cen, "censoring", ["use_followup_status", "followup_status_values", "event_types", "require_event_confirmed"])
    statuses = cen["followup_status_values"]
    if not isinstance(statuses, list) or any(s not in FOLLOWUP_STATUS_VALUES or s == "active" for s in statuses):
        _pfail("censoring.followup_status_values", f"只能取 {[s for s in FOLLOWUP_STATUS_VALUES if s != 'active']} 中的值")
    events = cen["event_types"]
    if not isinstance(events, list) or not all(isinstance(e, str) for e in events):
        _pfail("censoring.event_types", "应为事件类型列表，如 [\"esrd\", \"death\"]")
    bad = sorted({e for e in events if e not in EVENT_TYPES})
    if bad:
        _pfail("censoring.event_types", f"不认识的事件类型 {bad}；只能取登记系统事件表 events_long.event_type 的取值 {list(EVENT_TYPES)}"
                                        "（例如肾衰竭为 esrd，而不是 eskd）。拼写错误会让截尾静默失效，因此拒绝运行")
    p["censoring"] = {"use_followup_status": _pbool(cen["use_followup_status"], "censoring.use_followup_status"),
                      "followup_status_values": sorted(set(statuses)),
                      "event_types": sorted(set(events)),
                      "require_event_confirmed": _pbool(cen["require_event_confirmed"], "censoring.require_event_confirmed")}
    active = ([p["egfr_decline"]] if p["egfr_decline"] else []) + list(p["remission"].values())
    if any(not s["requires_pi_confirmation"] for s in active):
        c = p["confirmation"]
        if not c["confirmed_by"].strip() or pd.isna(date(c["confirmed_date"])):
            _pfail("confirmation", "已有章节标注 requires_pi_confirmation=false，但未填写确认人 confirmed_by 与确认日期 confirmed_date（YYYY-MM-DD）")
    return p


def load_protocol(path):
    """Returns (validated, raw, text, sha256, bytes); the run uses exactly these bytes."""
    path = Path(path)
    if path.is_symlink() or not path.is_file():
        raise DataContractError(f"研究方案文件不存在或不是普通文件：{path}")
    data = path.read_bytes()
    if len(data) > PROTOCOL_MAX_BYTES:
        raise DataContractError("研究方案文件过大（>1 MB）")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        raise DataContractError("研究方案文件必须是 UTF-8 编码的 JSON") from None

    def no_duplicates(pairs):
        keys = [k for k, _ in pairs]
        if len(keys) != len(set(keys)):
            raise DataContractError(f"研究方案文件含重复字段 {sorted({k for k in keys if keys.count(k) > 1})}")
        return dict(pairs)

    def no_constants(name):
        raise DataContractError(f"研究方案文件含非法数值 {name}")
    try:
        raw = json.loads(text.lstrip("﻿"), object_pairs_hook=no_duplicates, parse_constant=no_constants)
    except json.JSONDecodeError as exc:
        raise DataContractError(f"研究方案文件不是合法 JSON：第 {exc.lineno} 行第 {exc.colno} 列") from None
    return validate_protocol(raw), raw, text, hashlib.sha256(data).hexdigest(), data


def find_package_file(data_dir, name):
    """First regular file `name` in the input dir, its parent or grandparent
    (an export package keeps CSVs in analysis/data and metadata at the package root)."""
    data_dir = Path(data_dir)
    for folder in [data_dir, data_dir.parent, data_dir.parent.parent]:
        candidate = folder/name
        if candidate.is_file() and not candidate.is_symlink():
            return candidate
    return None


def declared_visit_upcr_units(data_dir, visits):
    """Explicit declarations of visit UPCR unit: per-row column, then export dictionary."""
    declared = []
    if "upcr_unit" in visits.columns:
        for value in sorted({v for v in visits["upcr_unit"] if v}):
            declared.append((value, "visits_long.upcr_unit"))
    candidate = find_package_file(data_dir, "data_dictionary.json")
    if candidate is not None:
        try:
            unit = json.loads(candidate.read_text(encoding="utf-8-sig")).get("visits", {}).get("upcr")
        except (ValueError, AttributeError):
            raise DataContractError(f"{candidate} 无法解析，不能确认随访 UPCR 单位") from None
        if isinstance(unit, str) and unit:
            declared.append((unit, str(candidate.name)))
    if not declared:
        declared.append((VISIT_UPCR_UNIT, "registry export contract (visits_long.upcr stored as mg/g)"))
    return declared


def check_protocol_units(protocol, data_dir, visits):
    declared = declared_visit_upcr_units(data_dir, visits)
    data_units = {UNIT_ALIASES["upcr"].get(u, u) for u, _ in declared}
    pu = protocol["units"]["upcr"]
    if len(data_units) != 1 or data_units != {pu} or pu != ANALYSIS_UNITS["upcr"]:
        detail = "；".join(f"{u}（来源 {src}）" for u, src in declared)
        raise DataContractError(f"研究方案 units.upcr={pu} 与数据显式单位不一致：{detail}。脚本不会自动换算阈值；"
                                f"请由研究方案负责人把阈值改写为数据单位 {VISIT_UPCR_UNIT} 并修改 units.upcr，或核对数据来源后重新导出。")
    if protocol["units"]["egfr"] != ANALYSIS_UNITS["egfr"]:
        raise DataContractError("研究方案 units.egfr 必须为 mL/min/1.73m2")
    if protocol["units"].get("scr", ANALYSIS_UNITS["scr"]) != ANALYSIS_UNITS["scr"]:
        raise DataContractError(f"研究方案 units.scr={protocol['units']['scr']} 与数据单位 umol/L 不一致；拒绝运行，不自动换算。")
    return [{"unit": u, "source": src} for u, src in declared]


def read_merge_manifest(data_dir):
    """MERGE_MANIFEST.json written by merge_centers.py, if the input is a merged dataset.

    Returns (summary for RUN_LOG or None, {(project_id, center_code): {table: [columns
    absent from that center's source export]}}).
    """
    path = Path(data_dir)/"MERGE_MANIFEST.json"
    if path.is_symlink() or not path.is_file():
        return None, {}
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
        centers = {c["source_dir"]: [tuple(pc) for pc in c["project_centers"]] for c in manifest.get("centers", [])}
        mismatch = list(manifest.get("column_mismatch", []))
        missing = {}
        for m in mismatch:
            for pc in centers.get(m["source_dir"], []):
                missing.setdefault(pc, {}).setdefault(m["table"], []).extend(m["missing_columns"])
    except (ValueError, KeyError, TypeError, AttributeError):
        raise DataContractError(f"{path} 无法解析；请重新运行 merge_centers.py 生成合并数据") from None
    summary = {"sha256": sha256(path), "status": manifest.get("status"), "allow_missing_columns": bool(manifest.get("allow_missing_columns", False)),
               "column_mismatch": mismatch, "input_sha256": manifest.get("input_sha256")}
    return summary, missing


def apply_consent_filter(frames, source_missing=None):
    """Default population: consent_status in consented/waiver when that column exists.

    Inclusion and exclusion are counted per project/center so a center whose export
    lacked the column, or whose consent was never recorded, is visible rather than
    silently lost.
    """
    source_missing = source_missing or {}
    baseline = frames["patients_baseline"]
    if "consent_status" not in baseline.columns:
        return frames, {"column_present": False, "applied": False, "excluded_n": 0, "by_center": [], "fully_excluded_centers": [],
                        "note": "未提供同意状态（patients_baseline 无 consent_status 列），未按知情同意筛选；请自行核对伦理与同意范围"}
    status = baseline["consent_status"].map(lambda v: str(v).strip())
    keep = status.isin(CONSENT_INCLUDED)
    labels = status.map(lambda s: s if s in CONSENT_VALUES else ("(empty)" if s == "" else "(unrecognised)"))
    counts = {k: int(v) for k, v in labels[~keep].value_counts().items()}
    by_center = []
    if not baseline.empty:
        tagged = baseline[["project_id", "center_code"]].assign(_keep=keep.values, _label=labels.values)
        for (pid, cc), sub in tagged.groupby(["project_id", "center_code"], sort=True):
            excluded = sub[~sub._keep]
            by_center.append({"project_id": pid, "center_code": cc, "patients_n": len(sub), "included_n": int(sub._keep.sum()),
                              "excluded_n": len(excluded),
                              "excluded_by_status": {k: int(v) for k, v in sorted(excluded._label.value_counts().items())},
                              "consent_status_column_absent_in_source_export": "consent_status" in source_missing.get((pid, cc), {}).get("patients_baseline", [])})
    excluded = baseline[~keep]
    drop = {tuple(r[k] for k in KEY) for r in excluded.to_dict("records")}
    out = {}
    for table, df in frames.items():
        mask = df.apply(lambda r: tuple(r[k] for k in KEY) not in drop, axis=1) if not df.empty else pd.Series([], dtype=bool)
        out[table] = df[mask].reset_index(drop=True) if not df.empty else df
    return out, {"column_present": True, "applied": True, "included_statuses": list(CONSENT_INCLUDED),
                 "patients_before_n": len(baseline), "included_n": int(keep.sum()), "excluded_n": int((~keep).sum()),
                 "excluded_by_status": dict(sorted(counts.items())), "by_center": by_center,
                 "fully_excluded_centers": [f"{c['project_id']}/{c['center_code']}" for c in by_center if c["included_n"] == 0],
                 "note": "默认只纳入 consent_status 为 consented 或 waiver 的患者；被排除患者的全部记录不进入任何输出。" + CONSENT_HINT_ZH}


def consent_failure_message(info):
    """Message when a project/center (or the whole input) has no included patient."""
    if not info["column_present"]:
        return None
    bad = [c for c in info["by_center"] if c["included_n"] == 0]
    if not bad and info["included_n"] > 0:
        return None
    lines = ["知情同意筛选后纳入人数为 0 的项目/中心如下，已停止分析，避免部分中心被整体排除后仍报告成功：" if bad
             else "知情同意筛选后纳入总人数为 0，已停止分析："]
    for c in bad:
        statuses = "，".join(f"{k} {v}" for k, v in c["excluded_by_status"].items())
        absent = "；该中心的源导出缺少 consent_status 列（合并时以空值填充）" if c["consent_status_column_absent_in_source_export"] else ""
        lines.append(f"  项目 {c['project_id']} / 中心 {c['center_code']}：{c['patients_n']} 名患者全部被排除（{statuses}）{absent}")
    lines.append(f"总计：纳入 {info['included_n']}，排除 {info['excluded_n']}（{json.dumps(info['excluded_by_status'], ensure_ascii=False)}）。")
    lines.append("原因：分析只纳入 consent_status 为 consented（已同意）或 waiver（伦理批准豁免）的患者。" + CONSENT_HINT_ZH)
    if "(empty)" in info["excluded_by_status"]:
        lines.append("空值（(empty)）通常表示该中心导出时数据库尚未启用知情同意字段，或合并时该中心导出缺少 consent_status 列。")
    lines.append("解决：在工作台为这些患者登记知情同意状态后重新导出；多中心合并请用同一版本工作台导出。若研究方案确实不分析这些中心，"
                 "可加 --allow-fully-excluded-centers 明确放行，运行日志会记录。")
    lines.append("English: after consent filtering, " + (f"{len(bad)} project/center group(s) have no included patient. " if bad else "no patient is included. ")
                 + CONSENT_HINT_EN + " To proceed deliberately pass --allow-fully-excluded-centers (recorded in RUN_LOG.json).")
    return "\n".join(lines)


def protocol_term_matches(protocol, frames):
    """How many records match each confirmed module / event type of the protocol.

    A confirmed module or censoring event type with zero matches is usually a spelling
    or version mismatch that would make a rule silently inert.
    """
    patients, events = frames["patients_baseline"], frames["events_long"]
    modules = patients.module.map(lambda m: str(m).strip().upper()).value_counts() if not patients.empty else pd.Series(dtype=int)
    ed = protocol["egfr_decline"]
    egfr_active = bool(ed) and not ed["requires_pi_confirmation"]
    remission = [m for m, s in protocol["remission"].items() if not s["requires_pi_confirmation"]]
    out = {"remission_modules": {m: int(modules.get(m, 0)) for m in remission},
           "egfr_decline_modules": {m: int(modules.get(m, 0)) for m in (ed["modules"] or [])} if egfr_active and ed["modules"] else {},
           "event_types": {}}
    if egfr_active or remission:
        types = events.event_type.map(lambda x: str(x).strip().lower()) if not events.empty else pd.Series(dtype=object)
        confirmed = events.confirmed.map(lambda x: str(x).strip().lower() in {"true", "t", "1"}) if not events.empty else pd.Series(dtype=bool)
        for t in protocol["censoring"]["event_types"]:
            hit = types == t
            out["event_types"][t] = {"records": int(hit.sum()), "confirmed": int((hit & confirmed).sum())}
    unmatched = ([f"remission.{m}（病种 {m} 无患者）" for m, n in out["remission_modules"].items() if n == 0]
                 + [f"egfr_decline.modules {m}（无患者）" for m, n in out["egfr_decline_modules"].items() if n == 0]
                 + [f"censoring.event_types {t}（事件表无此类型记录）" for t, c in out["event_types"].items() if c["records"] == 0])
    return out, unmatched


def unmatched_terms_message(unmatched):
    return ("研究方案中已确认使用的以下取值在本次数据中没有任何匹配：" + "；".join(unmatched) + "。"
            "这通常是拼写或版本不一致（例如病种 IGA 应为 IGAN、事件 eskd 应为 esrd），会让规则静默失效，因此拒绝运行。"
            "若核对后确认本数据中确实没有此类患者或事件，请加 --allow-unmatched-protocol-terms 明确放行，运行日志会记录匹配数。\n"
            "English: confirmed protocol terms with zero matches in the data: " + "; ".join(unmatched)
            + ". Check spelling against the registry vocabulary, or pass --allow-unmatched-protocol-terms to proceed deliberately.")


def _daily(sub, metric, unusable_col):
    """One row per date: mean of the usable values that day. value is NaN for a date
    whose only measurements were recorded but excluded (unverified source, invalid
    value): such a date is an UNKNOWN measurement, not an absent one."""
    cols = ["visit_date", "value", "days", "n_usable", "n_unusable"]
    if sub.empty:
        return pd.DataFrame(columns=cols)
    flag = sub[unusable_col].astype(bool) if unusable_col in sub else pd.Series(False, index=sub.index)
    s = sub[sub.visit_date.notna() & (sub[metric].notna() | flag)].assign(_unusable=flag.astype(int))
    if s.empty:
        return pd.DataFrame(columns=cols)
    return (s.groupby("visit_date", as_index=False)
              .agg(value=(metric, "mean"), days=("days_from_baseline", "first"), n_usable=(metric, "count"), n_unusable=("_unusable", "sum"))
              .sort_values("visit_date", kind="stable").reset_index(drop=True))


def _protocol_baseline(rule, record_value, daily):
    """Returns (value, source_label, followup_start_day). Follow-up is strictly after it.
    `daily` holds usable measurements only."""
    window = rule["window_days"]
    start_day = max(0, window[1]) if window else 0
    if rule["source"] in ("baseline_record", "baseline_record_else_visit_window") and np.isfinite(record_value):
        return record_value, "baseline_record", start_day
    if rule["source"] == "baseline_record":
        return np.nan, "baseline_record_unavailable", start_day
    cand = daily[daily.days.between(window[0], window[1])] if not daily.empty else daily
    if cand.empty:
        return np.nan, "no_measurement_in_baseline_window", start_day
    if rule["aggregate"] == "mean":
        return float(cand.value.mean()), f"visit_window_mean_n{len(cand)}", start_day
    cand = cand.assign(distance=cand.days.abs()).sort_values(["distance", "visit_date"], kind="stable")
    return float(cand.iloc[0].value), f"visit_window_nearest_day{int(cand.iloc[0].days)}", start_day


def _evaluate_runs(seq, conf, window):
    last = len(seq) - 1
    runs, cur = [], []
    for i, (_, _, q) in enumerate(seq):
        if q is not None and bool(q):
            cur.append(i)
        elif cur:
            runs.append(cur)
            cur = []
    if cur:
        runs.append(cur)
    detail = "never_met_criterion"
    for run in runs:
        interrupted = run[-1] < last and seq[run[-1] + 1][2] is None
        if window:
            # Only measurements inside the assessment window count; the run's first
            # in-window measurement must fall within it (confirmations may follow later).
            run = [i for i in run if seq[i][1] >= window[0]]
            if not run or seq[run[0]][1] > window[1]:
                detail = "met_only_outside_assessment_window"
                continue
        if conf["require_last_below"] and run[-1] != last:
            detail = "met_then_interrupted_by_unknown_measurement" if interrupted else "met_then_not_sustained_to_last_measurement"
            continue
        first = seq[run[0]][0]
        for n, j in enumerate(run, start=1):
            if n >= conf["min_measurements"] and (seq[j][0] - first).days >= conf["min_interval_days"]:
                return {"status": "met_per_protocol", "status_detail": "sustained_confirmed", "event_date": first,
                        "event_days_from_baseline": seq[run[0]][1], "confirmation_date": seq[j][0], "n_confirming_measurements": n}
        if len(run) >= conf["min_measurements"]:
            detail = "confirmation_interval_too_short"
        elif run[-1] == last:
            detail = "met_at_last_measurement_awaiting_confirmation"
        elif interrupted:
            detail = "not_confirmed_interrupted_by_unknown_measurement"
        else:
            detail = "transient_not_confirmed_recovered"
    return {"status": "not_met", "status_detail": detail}


def evaluate_sustained(seq, conf, window=None):
    """seq: date-ordered [(date, day, qualifies)], one per distinct date, already censored.

    qualifies is True, False, or None for an UNKNOWN measurement (recorded but
    excluded, or a required same-day value is missing). A run is consecutive
    qualifying measurements; a non-qualifying OR unknown measurement ends it, so an
    unknown is never bridged over and never counted as qualifying. The endpoint is met
    when a run has >=min_measurements measurements whose first and confirming dates are
    >=min_interval_days apart; with require_last_below the run must also reach the last
    measurement. With an assessment window, the event date is the first measurement of
    the qualifying run within the window. If the result would be met had the unknown
    measurements qualified, status_detail says unknown_measurement_prevents_confirmation.
    """
    if not seq:
        return {"status": "not_evaluable", "status_detail": "no_followup_measurements"}
    if all(q is None for _, _, q in seq):
        return {"status": "not_evaluable", "status_detail": "no_usable_followup_measurements"}
    result = _evaluate_runs(seq, conf, window)
    if result["status"] != "met_per_protocol" and any(q is None for _, _, q in seq):
        optimistic = _evaluate_runs([(d, day, True if q is None else q) for d, day, q in seq], conf, window)
        if optimistic["status"] == "met_per_protocol":
            return {"status": "not_met", "status_detail": "unknown_measurement_prevents_confirmation"}
    return result


def censoring_for(r, events, protocol, has_followup_status, issues):
    """Censoring candidates for one patient with a valid baseline date.

    Returns (items, reasons): items are (date, reason, kind) where kind is
    "kf_death_event" for a kidney failure (esrd/dialysis_start/kidney_transplant) or death event listed in censoring.event_types (and
    confirmed when required) and "other" otherwise; reasons lists, in order, every
    reason found, including ones whose date could not be used.
    """
    cen = protocol["censoring"]
    anchor = r["baseline_anchor"]
    items, reasons = [], []
    if cen["use_followup_status"] and has_followup_status:
        status = str(r.get("followup_status", "")).strip()
        if status in cen["followup_status_values"]:
            sd = date(r.get("followup_status_date", ""))
            if pd.isna(sd):
                issue(issues, "patients_baseline", r, "followup_status_date_missing",
                      f"followup_status={status} without valid followup_status_date; not used as censoring date")
                reasons.append(f"followup_status_{status}_date_missing")
            elif sd < anchor:
                issue(issues, "patients_baseline", r, "followup_status_date_before_baseline",
                      f"followup_status={status} dated {sd:%Y-%m-%d}, before baseline {anchor:%Y-%m-%d}; date ignored for censoring "
                      "(baseline and follow-up measurements kept); correct the date in the workbench")
                reasons.append(f"followup_status_{status}_date_before_baseline_ignored")
            else:
                items.append((sd, f"followup_status_{status}", "other"))
                reasons.append(f"followup_status_{status}")
    sub = events
    for k in KEY:
        sub = sub[sub[k] == r[k]]
    for e in sub.to_dict("records"):
        etype = str(e["event_type"]).strip().lower()
        if etype not in cen["event_types"]:
            continue
        if cen["require_event_confirmed"] and str(e["confirmed"]).strip().lower() not in {"true", "t", "1"}:
            continue
        ed = date(e["event_date"])
        if pd.isna(ed):
            issue(issues, "events_long", e, "censoring_event_date_missing", f"{etype} event without valid event_date; not used for censoring")
            reasons.append(f"event_{etype}_date_missing")
        elif ed <= anchor:
            issue(issues, "events_long", e, "censoring_event_not_after_baseline",
                  f"{etype} event dated {ed:%Y-%m-%d}, not after baseline {anchor:%Y-%m-%d}; not used for censoring")
            reasons.append(f"event_{etype}_not_after_baseline_ignored")
        else:
            items.append((ed, f"event_{etype}", "kf_death_event" if etype in KF_DEATH_EVENT_TYPES else "other"))
            reasons.append(f"event_{etype}")
    return items, reasons


def _egfr_source_counts(rows):
    """Composition of eGFR measurements used: recomputed from creatinine, manual value
    (stored value tagged manual), excluded (recorded but unverified/unusable)."""
    if rows.empty:
        return ""
    counts = {"recomputed": int(rows.egfr_analysis_source.isin(RECOMPUTED_EGFR_SOURCES).sum()),
              "manual": int((rows.egfr_analysis_source == "recorded_manual").sum()),
              "excluded": int(rows.egfr_unusable.astype(bool).sum())}
    return ";".join(f"{k}:{v}" for k, v in counts.items() if v)


def _remission_eligibility(spec, upcr_base, upcr_src):
    """(eligible, detail). Remission is only assessed when baseline UPCR is known and
    not already in the remission range (or >= baseline_upcr_at_least when set)."""
    if not (np.isfinite(upcr_base) and upcr_base > 0):
        return False, "baseline_upcr_unavailable:" + upcr_src
    at_least = spec["baseline_upcr_at_least"]
    if at_least is not None:
        if _q(upcr_base) >= _q(at_least):
            return True, ""
        return False, f"baseline_upcr_below_eligibility:{upcr_base:g}<{at_least:g}(baseline_upcr_at_least)"
    for level in ("complete", "partial"):
        cut = spec[level]["upcr_below"] if spec[level] else None
        if cut is not None:
            if _q(upcr_base) > _q(cut):
                return True, ""
            return False, f"baseline_upcr_not_above_remission_threshold:{upcr_base:g}<={cut:g}({level}.upcr_below)"
    return True, ""


def protocol_endpoints(protocol, patients, visits, events, issues):
    tag = f"{protocol['protocol_id']}@{protocol['version']}"
    determination = f"已按方案 {tag} 判定"
    has_fs = "followup_status" in patients.columns
    rows, tps, skipped = [], [], {}
    ed = protocol["egfr_decline"]
    egfr_active = bool(ed) and not ed["requires_pi_confirmation"]
    kf_mode = ed["kidney_failure_or_death_before_confirmation"] if ed else None
    if ed and ed["requires_pi_confirmation"]:
        skipped["egfr_decline"] = "requires_pi_confirmation=true"
    for module, spec in protocol["remission"].items():
        if spec["requires_pi_confirmation"]:
            skipped[f"remission.{module}"] = "requires_pi_confirmation=true"
    present_modules = sorted({str(m).upper() for m in patients.module})
    not_covered = {m: int((patients.module.str.upper() == m).sum()) for m in present_modules if m not in protocol["remission"]}
    for section, why in skipped.items():
        issues.append({"table": "protocol", **{k: "" for k in KEY}, "record_id": "", "code": "protocol_section_not_confirmed",
                       "detail": f"{section} not computed: protocol marks it {why}; PI must confirm definitions first"})
    for module, n in not_covered.items():
        issues.append({"table": "protocol", **{k: "" for k in KEY}, "record_id": "", "code": "protocol_module_not_covered",
                       "detail": f"Remission not computed for module {module} ({n} patients): protocol {tag} has no definition for it"})
    if ed and egfr_active and ed["modules"]:
        for module in present_modules:
            if module not in ed["modules"]:
                issues.append({"table": "protocol", **{k: "" for k in KEY}, "record_id": "", "code": "protocol_module_not_covered",
                               "detail": f"eGFR decline not computed for module {module}: excluded by egfr_decline.modules"})
    upcr_rule = protocol["baseline"]["upcr"]
    peak_start = upcr_rule["window_days"][0] if upcr_rule["window_days"] else 0
    for r in patients.to_dict("records"):
        sub = visits
        for k in KEY:
            sub = sub[sub[k] == r[k]]
        module = str(r["module"]).upper()
        common = {**{k: r[k] for k in KEY}, "module": r["module"], "protocol": tag, "determination": determination}
        anchor = r["baseline_anchor"]
        items, reasons = censoring_for(r, events, protocol, has_fs, issues) if pd.notna(anchor) else ([], [])
        censor_date = min(d for d, _, _ in items) if items else pd.NaT
        censor_reason = ";".join(reasons)
        kf_dates = [d for d, _, kind in items if kind == "kf_death_event"]
        kf_first = min(kf_dates) if kf_dates else pd.NaT
        # The KF/death event only matters for count_as_event when nothing censored earlier.
        kf_before_other = pd.notna(kf_first) and kf_first == censor_date
        kf_basis = ";".join(sorted({reason for d, reason, kind in items if kind == "kf_death_event" and d == kf_first}))
        cens = sub if pd.isna(censor_date) else sub[sub.visit_date <= censor_date]
        egfr_daily, upcr_daily = _daily(cens, "egfr", "egfr_unusable"), _daily(cens, "upcr", "upcr_unusable")
        egfr_known = egfr_daily[egfr_daily.value.notna()] if not egfr_daily.empty else egfr_daily
        upcr_known = upcr_daily[upcr_daily.value.notna()] if not upcr_daily.empty else upcr_daily
        egfr_base, egfr_src, egfr_start = _protocol_baseline(protocol["baseline"]["egfr"], number(r["baseline_egfr"]), egfr_known)
        upcr_base, upcr_src, upcr_start = _protocol_baseline(upcr_rule, number(r["baseline_upcr_final"]), upcr_known)
        egfr_fu_all = egfr_daily[egfr_daily.days > egfr_start] if not egfr_daily.empty else egfr_daily
        upcr_fu_all = upcr_daily[upcr_daily.days > upcr_start] if not upcr_daily.empty else upcr_daily
        egfr_fu = egfr_fu_all[egfr_fu_all.value.notna()] if not egfr_fu_all.empty else egfr_fu_all
        upcr_fu = upcr_fu_all[upcr_fu_all.value.notna()] if not upcr_fu_all.empty else upcr_fu_all
        egfr_by_date = dict(zip(egfr_known.visit_date, egfr_known.value)) if not egfr_known.empty else {}
        egfr_unknown_dates = set(egfr_daily[egfr_daily.value.isna()].visit_date) if not egfr_daily.empty else set()
        dated = cens[cens.visit_date.notna()]

        def emit(endpoint, result, base, unit, src, fu, extra="", unknown=None, egfr_sources=""):
            last_date = fu.visit_date.iloc[-1] if len(fu) else pd.NaT
            unknown = {k: v for k, v in (unknown or {}).items() if v}
            rows.append({**common, "endpoint": endpoint, "status": result["status"], "status_detail": result["status_detail"],
                         "event_basis": result.get("event_basis", ""),
                         "event_date": result.get("event_date", pd.NaT), "confirmation_date": result.get("confirmation_date", pd.NaT),
                         "event_days_from_baseline": result.get("event_days_from_baseline", np.nan),
                         "n_confirming_measurements": result.get("n_confirming_measurements", np.nan),
                         "baseline_value": base, "baseline_unit": unit, "baseline_source": src, "n_followup_measurements": len(fu),
                         "n_unknown_measurements": sum(unknown.values()),
                         "unknown_measurement_reasons": ";".join(f"{k}:{v}" for k, v in sorted(unknown.items())),
                         "egfr_source_counts": egfr_sources,
                         "last_measurement_date": last_date, "last_measurement_days_from_baseline": fu.days.iloc[-1] if len(fu) else np.nan,
                         "censor_date": censor_date, "censor_reason": censor_reason, "criteria_not_evaluated": extra})
            if result["status"] == "not_evaluable":
                issue(issues, "endpoints_per_protocol", r, "endpoint_not_evaluable", f"{endpoint}: {result['status_detail']}")

        if egfr_active and (ed["modules"] is None or module in ed["modules"]):
            egfr_rows = dated[dated.days_from_baseline > egfr_start]
            for pct in ed["thresholds_pct_decline"]:
                name = f"egfr_decline_{pct:g}pct_sustained" + ("_or_kidney_failure_death" if kf_mode == "count_as_event" else "")
                unknown = {}
                if pd.isna(anchor):
                    result = {"status": "not_evaluable", "status_detail": "baseline_date_unavailable"}
                elif not (np.isfinite(egfr_base) and egfr_base > 0):
                    result = {"status": "not_evaluable", "status_detail": "baseline_egfr_unavailable:"+egfr_src}
                else:
                    remaining = 100 - _q(pct)
                    seq = [(d, int(day), None if pd.isna(v) else le_pct_of(v, egfr_base, remaining))
                           for d, day, v in zip(egfr_fu_all.visit_date, egfr_fu_all.days, egfr_fu_all.value)]
                    unknown = {"egfr_unusable": sum(q is None for _, _, q in seq)}
                    result = evaluate_sustained(seq, ed["confirmation"]) if seq else {"status": "not_evaluable", "status_detail": "no_followup_measurements"}
                    if result["status"] == "met_per_protocol":
                        result["event_basis"] = "egfr_measurements"
                    elif kf_mode == "count_as_event" and kf_before_other:
                        result = {"status": "met_per_protocol", "status_detail": "kidney_failure_or_death_before_confirmation",
                                  "event_basis": kf_basis, "event_date": kf_first, "event_days_from_baseline": int((kf_first - anchor).days)}
                emit(name, result, egfr_base, ANALYSIS_UNITS["egfr"], egfr_src, egfr_fu, unknown=unknown, egfr_sources=_egfr_source_counts(egfr_rows))
        spec = protocol["remission"].get(module)
        if spec and not spec["requires_pi_confirmation"]:
            extra = "; ".join(spec["additional_criteria_not_computed"])
            eligible, eligibility_detail = _remission_eligibility(spec, upcr_base, upcr_src)
            peak_pool = upcr_known[upcr_known.days >= peak_start] if not upcr_known.empty else upcr_known
            for level in ("complete", "partial"):
                crit = spec[level]
                if crit is None:
                    continue
                name = f"remission_{level}_sustained"
                need_reduction = crit["upcr_reduction_pct_at_least"] is not None
                need_egfr = crit["egfr_stable_max_decline_pct"] is not None
                unknown, egfr_sources = {}, ""
                if pd.isna(anchor):
                    result = {"status": "not_evaluable", "status_detail": "baseline_date_unavailable"}
                elif not eligible:
                    result = {"status": "not_evaluable", "status_detail": eligibility_detail}
                elif need_egfr and not (np.isfinite(egfr_base) and egfr_base > 0):
                    result = {"status": "not_evaluable", "status_detail": "baseline_egfr_unavailable_for_stability:"+egfr_src}
                else:
                    seq = []
                    unknown = {"upcr_unusable": 0, "no_same_day_egfr": 0, "same_day_egfr_unusable": 0}
                    for d, day, v in zip(upcr_fu_all.visit_date, upcr_fu_all.days, upcr_fu_all.value):
                        if pd.isna(v):
                            unknown["upcr_unusable"] += 1
                            seq.append((d, int(day), None))
                            continue
                        ok = True
                        if crit["upcr_below"] is not None:
                            ok = ok and _q(v) < _q(crit["upcr_below"])
                        if need_reduction:
                            ref = upcr_base
                            if crit["upcr_reduction_reference"] == "peak":
                                earlier = peak_pool[peak_pool.days < day].value if not peak_pool.empty else []
                                if len(earlier):
                                    ref = max(ref, float(max(earlier)))
                            ok = ok and le_pct_of(v, ref, 100 - _q(crit["upcr_reduction_pct_at_least"]))
                        q = bool(ok)
                        if ok and need_egfr:
                            if d in egfr_by_date:
                                q = ge_pct_of(egfr_by_date[d], egfr_base, 100 - _q(crit["egfr_stable_max_decline_pct"]))
                            else:
                                # Unknown: never counted as qualifying, and it interrupts a run (same rule as an excluded eGFR).
                                unknown["same_day_egfr_unusable" if d in egfr_unknown_dates else "no_same_day_egfr"] += 1
                                q = None
                        seq.append((d, int(day), q))
                    result = evaluate_sustained(seq, spec["confirmation"], spec["assessment_window_days"]) if seq else {"status": "not_evaluable", "status_detail": "no_followup_measurements"}
                    if result["status"] == "met_per_protocol":
                        result["event_basis"] = "upcr_measurements"
                        if extra:
                            result = {**result, "status": "upcr_egfr_criteria_met_other_criteria_not_evaluated"}
                    if need_egfr:
                        egfr_sources = _egfr_source_counts(dated[dated.visit_date.isin(set(upcr_fu_all.visit_date))])
                emit(name, result, upcr_base, ANALYSIS_UNITS["upcr"], upcr_src, upcr_fu, extra, unknown=unknown, egfr_sources=egfr_sources)
        for nominal in protocol["timepoints_days"]:
            w = protocol["visit_window_days"]
            row = {**common, "nominal_day": nominal, "window_days": w, "status": "no_visit_in_window", "visit_id": "",
                   "visit_date": pd.NaT, "days_from_baseline": np.nan, "egfr": np.nan, "upcr": np.nan,
                   "egfr_baseline_protocol": egfr_base, "upcr_baseline_protocol": upcr_base,
                   "egfr_pct_change": np.nan, "upcr_pct_change": np.nan, "tie_count": 0}
            if pd.isna(anchor):
                row["status"] = "baseline_date_unavailable"
            else:
                cand = cens[cens.days_from_baseline.between(max(1, nominal - w), nominal + w)].copy()
                if not cand.empty:
                    cand["distance"] = (cand.days_from_baseline - nominal).abs()
                    cand = cand.sort_values(["distance", "visit_date", "id"], kind="stable")
                    pk = cand.iloc[0]
                    row.update(status="visit_selected", visit_id=pk["id"], visit_date=pk.visit_date, days_from_baseline=pk.days_from_baseline,
                               egfr=pk.egfr, upcr=pk.upcr, tie_count=int((cand.distance == pk.distance).sum()),
                               egfr_pct_change=100*(pk.egfr-egfr_base)/egfr_base if np.isfinite(egfr_base) and egfr_base > 0 else np.nan,
                               upcr_pct_change=100*(pk.upcr-upcr_base)/upcr_base if np.isfinite(upcr_base) and upcr_base > 0 else np.nan)
            tps.append(row)
    computed = (["egfr_decline"] if egfr_active else []) + [f"remission.{m}" for m, s in protocol["remission"].items() if not s["requires_pi_confirmation"]]
    summary = {"tag": tag, "sections_computed": computed, "sections_skipped": skipped, "modules_not_covered_for_remission": not_covered,
               "followup_status_column_present": has_fs, "egfr_decline_kidney_failure_or_death_before_confirmation": kf_mode}
    endpoints = pd.DataFrame(rows, columns=ENDPOINT_COLUMNS)
    for col in ["event_days_from_baseline", "n_confirming_measurements", "last_measurement_days_from_baseline", "n_unknown_measurements"]:
        endpoints[col] = pd.to_numeric(endpoints[col]).astype("Int64")  # whole days/counts, blank when absent
    return endpoints, pd.DataFrame(tps, columns=TIMEPOINT_COLUMNS), summary


def egfr_source_summary(visits):
    """RUN_LOG: where every analysed eGFR came from."""
    dated = visits[visits.visit_date.notna()]
    recomputed = dated.egfr_analysis_source.isin(RECOMPUTED_EGFR_SOURCES)
    stored = dated.egfr_recorded.map(number)
    diff = (dated.egfr - stored).abs()
    return {"dated_visits_n": len(dated),
            "recomputed_from_creatinine_n": int(recomputed.sum()),
            "recomputed_stored_value_agrees_n": int((recomputed & (diff <= 0.11)).sum()),
            "recomputed_stored_value_differs_n": int((recomputed & (diff > 0.11)).sum()),
            "recomputed_no_stored_value_n": int((recomputed & stored.isna()).sum()),
            "recomputed_schwartz_2009_n": int((dated.egfr_analysis_source == RECOMPUTED_EGFR_PAEDIATRIC).sum()),
            "manual_value_used_n": int((dated.egfr_analysis_source == "recorded_manual").sum()),
            "excluded_recorded_but_unusable_n": int(dated.egfr_unusable.astype(bool).sum()),
            "no_egfr_or_creatinine_recorded_n": int((~dated.egfr_unusable.astype(bool) & dated.egfr.isna()).sum()),
            "note": "recomputed=按 CKD-EPI 2021 由血肌酐重算，1–17 岁按本次身高用床旁 Schwartz 2009 重算（recomputed_schwartz_2009_n；数据库存储值只用于核对）；manual=数据库中标记为 manual 的原值；"
                    "excluded=有记录但来源未核实或无法使用，未进入分析，在按方案判定的序列中记为未知测量"}


def provenance_files(data_dir, input_hashes):
    """RUN_LOG: hashes and identifiers of the export package files next to the input."""
    data_dir = Path(data_dir)
    out = {}
    for name in ("data_dictionary.json", "export_manifest.json", "research_data.json"):
        path = find_package_file(data_dir, name)
        if path is None:
            out[name] = None
            continue
        entry = {"location": os.path.relpath(path, data_dir).replace(os.sep, "/"), "sha256": sha256(path), "bytes": path.stat().st_size}
        if name == "export_manifest.json":
            try:
                manifest = json.loads(path.read_text(encoding="utf-8-sig"))
            except ValueError:
                manifest = None
                entry["parse_error"] = True
            if isinstance(manifest, dict):
                entry["identifiers"] = {k: manifest.get(k) for k in ("format", "project_id", "center_code", "exported_at", "metadata_register_id", "source_content_sha256")}
                listed = manifest.get("files") if isinstance(manifest.get("files"), dict) else {}
                prefix = os.path.relpath(data_dir, path.parent).replace(os.sep, "/")
                check = {}
                for csv_name, digest in input_hashes.items():
                    spec = listed.get(csv_name if prefix == "." else f"{prefix}/{csv_name}")
                    if isinstance(spec, dict) and "sha256" in spec:
                        check[csv_name] = "matches_export_manifest" if spec["sha256"] == digest else "differs_from_export_manifest"
                    else:
                        check[csv_name] = "not_listed_in_export_manifest"
                entry["input_csv_check"] = check
        out[name] = entry
    snapshots = sorted({p for folder in [data_dir, data_dir.parent, data_dir.parent.parent]
                        for p in folder.glob("registry_frozen_*.json") if p.is_file() and not p.is_symlink()})
    out["frozen_snapshots"] = [{"location": os.path.relpath(p, data_dir).replace(os.sep, "/"), "sha256": sha256(p)} for p in snapshots]
    return out


def manual_events(events, patients):
    out = events.merge(patients[KEY+["baseline_anchor"]], on=KEY, how="left", validate="many_to_one")
    out["analysis_status"] = pd.Series(index=out.index, dtype=object)
    for idx,r in out.iterrows():
        ed=date(r.event_date)
        confirmed=str(r.confirmed).strip().lower() in {"true","t","1"}
        if not confirmed: status="not_confirmed"
        elif pd.isna(ed): status="event_date_unavailable"
        elif pd.isna(r.baseline_anchor): status="baseline_date_unavailable"
        elif ed <= r.baseline_anchor: status="baseline_or_prebaseline_event"
        else: status="confirmed_in_followup"
        out.at[idx,"analysis_status"]=status
    return out


def patient_slopes(patients, visits):
    rows=[]
    for r in patients.to_dict("records"):
        sub=visits
        for k in KEY: sub=sub[sub[k]==r[k]]
        sub=sub[(sub.days_from_baseline>=0)&sub.egfr.notna()].copy()
        # Multiple observations per date get one date mean, explicitly reported.
        daily=sub.groupby("visit_date",as_index=False).agg(time_yr=("time_yr","first"),egfr=("egfr","mean"))
        row={**{k:r[k] for k in KEY},"status":"insufficient_distinct_dates","slope_mL_yr":np.nan,"follow_up_yr":0.0,"n_visits_used":len(sub),"n_distinct_dates":len(daily)}
        if len(daily)>=2:
            x=daily.time_yr.to_numpy(dtype=float);y=daily.egfr.to_numpy(dtype=float)
            dx=x-x.mean();denom=float(np.dot(dx,dx))
            if denom>0 and np.isfinite(denom):
                slope=float(np.dot(dx,y-y.mean())/denom)
                if np.isfinite(slope): row.update(status="descriptive_ols",slope_mL_yr=slope,follow_up_yr=float(x.max()-x.min()))
        rows.append(row)
    return pd.DataFrame(rows,columns=KEY+["status","slope_mL_yr","follow_up_yr","n_visits_used","n_distinct_dates"])


def trend_summary(visits, metric):
    df=visits[(visits.days_from_baseline>=0)&visits[metric].notna()].copy()
    if df.empty: return pd.DataFrame(columns=["month","mean","n_patients","sd_between_patient_month_means"])
    df["month"]=(df.days_from_baseline/30.4375).round().astype(int)
    per_patient=df.groupby(KEY+["month"],as_index=False)[metric].mean()
    return per_patient.groupby("month",as_index=False).agg(mean=(metric,"mean"),n_patients=(metric,"count"),sd_between_patient_month_means=(metric,"std"))


def table_one(patients, visits, meds, consent_info=None):
    """Descriptive baseline table only; never endpoint, remission or decline rates."""
    n=len(patients)
    rows=[{"Variable":"Patients","n":n,"denominator":n,"Value":n}, {"Variable":"Project-center groups","n":len(patients[["project_id","center_code"]].drop_duplicates()),"denominator":"","Value":""}]
    for sex in ["F","M"]:
        count=int((patients.sex.str.strip().str.upper()==sex).sum())
        rows.append({"Variable":f"Sex {sex}","n":count,"denominator":n,"Value":100*count/n if n else np.nan})
    for field in ["age_baseline","baseline_egfr","baseline_upcr_final"]:
        s=pd.to_numeric(patients[field],errors="coerce").dropna()
        rows.append({"Variable":field+" (median)","n":len(s),"denominator":n,"Value":s.median() if len(s) else np.nan,"Notes":f"IQR {s.quantile(.25):g}–{s.quantile(.75):g}" if len(s) else "not evaluable"})
        rows.append({"Variable":field+" (mean)","n":len(s),"denominator":n,"Value":s.mean() if len(s) else np.nan,"Notes":f"SD {s.std():g}" if len(s)>1 else "SD not evaluable"})
    for field,levels in [("oxford_m",[0,1]),("oxford_e",[0,1]),("oxford_s",[0,1]),("oxford_t",[0,1,2]),("oxford_c",[0,1,2])]:
        if field not in patients: continue
        applicable=patients[patients.module.str.upper()=="IGAN"]
        values=pd.to_numeric(applicable[field],errors="coerce")
        valid=values[values.isin(levels)]
        for level in levels:
            count=int((valid==level).sum())
            rows.append({"Variable":f"{field}={level}","n":count,"denominator":len(valid),"Value":100*count/len(valid) if len(valid) else np.nan,"Notes":f"IgAN applicable n={len(applicable)}; missing/invalid n={len(applicable)-len(valid)}"})
    if not meds.empty:
        any_n=len(meds[KEY].drop_duplicates())
        rows.append({"Variable":"Any medication recorded","n":any_n,"denominator":n,"Value":100*any_n/n if n else np.nan})
        for cls,sub in meds[meds.drug_class!=""].groupby("drug_class"):
            count=len(sub[KEY].drop_duplicates())
            rows.append({"Variable":"Medication class: "+cls,"n":count,"denominator":n,"Value":100*count/n if n else np.nan})
    if consent_info and consent_info.get("column_present"):
        rows.append({"Variable":"Excluded before analysis: consent not consented/waiver","n":consent_info["excluded_n"],"denominator":consent_info["patients_before_n"],"Value":consent_info["excluded_n"],"Notes":"counts by status in RUN_LOG.json"})
    return pd.DataFrame(rows)


def write_xlsx(path, sheets):
    # Literal strings in Excel, including strings beginning '='; no formula execution.
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active)
    for name,df in sheets.items():
        ws=wb.create_sheet(name[:31])
        for values in [list(df.columns)]+list(df.itertuples(index=False,name=None)):
            safe=[None if pd.isna(v) else (v.to_pydatetime() if isinstance(v,pd.Timestamp) else v) for v in values]
            ws.append(safe)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str): cell.data_type="s"
    wb.save(path)


def plot_descriptive_trend(df, metric, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=140)
    if df.empty:
        ax.text(0.5, 0.5, "No evaluable dated observations", ha="center", va="center", transform=ax.transAxes)
    else:
        ax.plot(df["month"], df["mean"], marker="o")
    ax.set(xlabel="Months from baseline", ylabel="eGFR (mL/min/1.73m²)" if metric=="egfr" else "UPCR (mg/g)", title="Descriptive mean: equal weight per patient-month; no confidence interval")
    ax.grid(alpha=0.2)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def actual_methods(log, protocol_text=None):
    consent = log["consent_filter"]
    if consent["column_present"]:
        consent_text = (f"Population: only patients with consent_status consented or waiver were analysed "
                        f"({consent['included_n']} included, {consent['excluded_n']} excluded; counts by status and by\n"
                        "project/center in RUN_LOG.json and qc_issues.csv). The registry default 'unconfirmed' and blank\n"
                        "values were excluded. All records of excluded patients were removed before any computation.")
        if consent.get("fully_excluded_centers"):
            consent_text += ("\nProject/center groups with no included patient (analysis explicitly allowed to proceed with\n"
                             f"--allow-fully-excluded-centers): {', '.join(consent['fully_excluded_centers'])}.")
    else:
        consent_text = ("Population: the input had no consent_status column, so no consent-based selection was\n"
                        "applied; the investigator must verify consent/ethics scope separately.")
    common = f"""# Methods actually executed — {VERSION}

Input: {log['patients_n']} patients and {log['visits_n']} visits, keyed by project_id,
center_code and the literal patient_code. Counts and file hashes are in RUN_LOG.json.
{consent_text}

Dates were parsed as ISO calendar dates; no missing baseline date or baseline unit
was imputed from subsequent visits or numeric magnitude. Unknown sex was not male.
Baseline eGFR and visits explicitly tagged CKD-EPI-2021-Cr used the adult creatinine
equation from Inker 2021 Table 2; serum creatinine conversion was 88.4 μmol/L per
mg/dL. Age was approximated by calendar year minus birth year (not exact birthday).
Manual eGFR remained tagged manual; values of unverified origin were excluded.
Reference: {EGFR_REFERENCE}

Annual review selected one whole visit within days 270–450, minimizing distance
to day 365, then earliest date, then lexical stable record ID. Ties are counted.
Missing values on that row remain missing. No annual visit means not evaluable.
This window is a descriptive preset and requires protocol review before inference.
"""
    if protocol_text is None:
        endpoint_text = """
No study protocol file was supplied. Unconfirmed eGFR candidates flag any single
post-baseline value <=60% or <=43% of baseline; unconfirmed IgAN proteinuria
candidates flag any single UPCR<300 mg/g or <=50% of baseline and <1000 mg/g.
These fixed thresholds are descriptive screening flags only: no persistence or
confirmation was required, they are not sustained, adjudicated or validated
endpoints or remission definitions, and no rates were derived from them.
Manual events remain individually retained; only explicitly confirmed, dated,
post-baseline events are labelled in followup.
"""
    else:
        p = log["protocol"]
        endpoint_text = f"""
Endpoints were determined by this script from study protocol {p['tag']}
(file {p['file_name']}, SHA-256 {p['sha256']}), confirmed by
{p['confirmed_by'] or '(not stated)'} on {p['confirmed_date'] or '(not stated)'} according to the file.
This is algorithmic determination from the protocol parameters, not manual clinical
adjudication. Sections computed: {', '.join(p['sections_computed']) or 'none'}.
Sections not computed: {json.dumps(p['sections_skipped'], ensure_ascii=False) or '{}'};
modules without a remission definition (not computed): {json.dumps(p['modules_not_covered_for_remission'], ensure_ascii=False)}.

Repeated measurements on one date were averaged to one measurement. Follow-up
measurements were those strictly after the protocol baseline (and after the end of
any baseline window). A measurement that was recorded but excluded (eGFR of
unverified source or without usable inputs, invalid UPCR), and a UPCR date without a
usable same-day eGFR when eGFR stability was required, was an UNKNOWN measurement: it
was never counted as qualifying and it interrupted a run (no bridging across it).
A criterion was met only by a run of consecutive qualifying measurements (no
non-qualifying or unknown measurement between) containing at least min_measurements
values whose first and confirming dates were at least min_interval_days apart; with
require_last_below the run also had to include the last measurement. The event date
is the first measurement of the qualifying run; when an assessment window was set,
only measurements on or after the window start counted and the event date (the first
in-window measurement of the run) had to fall within the window. Results that would
have been met had the unknown measurements qualified are labelled
unknown_measurement_prevents_confirmation; unknown counts are on each endpoint row.
Threshold comparisons were exact (rational arithmetic on the recorded decimal values).
Remission was assessed only when baseline UPCR was available and at least
baseline_upcr_at_least (or, when not set, above the remission upcr_below threshold);
otherwise it was not evaluable. A UPCR reduction was computed from the protocol
baseline value or, when the protocol states reference "peak", from the highest UPCR
between the baseline window start and the measurement being assessed.
All endpoints were anchored at the registry baseline_date (not treatment start).
Measurements were used up to and including the censoring date (protocol censoring
rules: follow-up status date and/or confirmed events; a follow-up status date before
baseline was ignored and reported); later measurements were not used. Handling of confirmed kidney failure (esrd, dialysis_start,
kidney_transplant) or death
events before a decline was confirmed: {p['egfr_decline_kidney_failure_or_death_before_confirmation'] or 'not applicable'}
(count_as_event: composite endpoint named *_or_kidney_failure_death with the event
date of that event; censor: censored at that date). UPCR thresholds were applied in
{ANALYSIS_UNITS['upcr']}, matching the explicit data unit; no threshold was rescaled.
Criteria the protocol lists as not computable (e.g. serum albumin) were not
evaluated and such results are labelled accordingly.

Full protocol text used for this run:

```json
{protocol_text.rstrip()}
```
"""
    tail = f"""
Individual descriptive OLS slopes were produced for {log['ols_evaluable_n']}
patients with at least two distinct dates. Same-date eGFR values were first averaged.
Month summaries first average within patient/month, then equally across patients.
No confidence intervals, cohort mixed-effects model, causal comparisons, survival
analysis, KFRE risk predictions or manual clinical endpoint adjudication were executed.
Missing values and exclusion reasons are recorded in qc_issues.csv. No imputation.
"""
    return common + endpoint_text + tail


def analyze(data_dir, output_dir, protocol_path=None, *, allow_fully_excluded_centers=False, allow_unmatched_protocol_terms=False):
    if Path(output_dir).is_symlink():
        raise DataContractError("Output must not be a symlink")
    data_dir,output_dir=Path(data_dir).resolve(),Path(output_dir).resolve()
    if output_dir.is_symlink() or (output_dir.exists() and any(p.name!=".keep" for p in output_dir.iterdir())):
        raise DataContractError("Output directory must be new/empty; use a new run folder to avoid stale results")
    protocol=protocol_raw_text=protocol_sha=protocol_bytes=None
    if protocol_path is not None:
        protocol,_,protocol_raw_text,protocol_sha,protocol_bytes=load_protocol(protocol_path)
    input_hashes={p.name:sha256(p) for p in sorted(data_dir.glob("*.csv"))}
    merge_info,source_missing=read_merge_manifest(data_dir)
    frames=load_tables(data_dir);issues=[]
    unit_declarations=check_protocol_units(protocol,data_dir,frames["visits_long"]) if protocol else None
    frames,consent_info=apply_consent_filter(frames,source_missing)
    failure=consent_failure_message(consent_info)
    if failure and not allow_fully_excluded_centers:
        raise DataContractError(failure)
    consent_info["proceeded_with_allow_fully_excluded_centers"]=bool(failure)
    for c in consent_info["by_center"]:
        issues.append({"table":"patients_baseline","project_id":c["project_id"],"center_code":c["center_code"],"patient_code":"","record_id":"",
                       "code":"consent_filter_center_counts" if c["included_n"] else "consent_filter_center_fully_excluded",
                       "detail":f"included {c['included_n']}/{c['patients_n']}; excluded {c['excluded_n']} by consent_status {json.dumps(c['excluded_by_status'],ensure_ascii=False)}"
                                +("; consent_status column absent in this center's source export (merged with --allow-missing-columns)" if c["consent_status_column_absent_in_source_export"] else "")
                                +("" if c["included_n"] else "; whole center excluded: record consent (consented/waiver) in the workbench; default 'unconfirmed' is excluded")})
    term_matches=None
    if protocol:
        term_matches,unmatched=protocol_term_matches(protocol,frames)
        if unmatched and not allow_unmatched_protocol_terms:
            raise DataContractError(unmatched_terms_message(unmatched))
        term_matches["unmatched"]=unmatched
        term_matches["proceeded_with_allow_unmatched_protocol_terms"]=bool(unmatched)
    patients,visits=derive(frames,issues)
    labs=normalize_labs(frames["labs_long"],issues)
    outcomes=choose_12month(patients,visits)
    events=manual_events(frames["events_long"],patients)
    slopes=patient_slopes(patients,visits)
    if protocol:
        endpoints,timepoints,psummary=protocol_endpoints(protocol,patients,visits,frames["events_long"],issues)
    else:
        declines,protein=candidates(patients,visits)
    qc=pd.DataFrame(issues,columns=["table"]+KEY+["record_id","code","detail"])
    if input_hashes != {p.name:sha256(p) for p in sorted(data_dir.glob("*.csv"))}:
        raise DataContractError("Input CSVs changed while reading; freeze the input files and retry")
    if protocol_path is not None and (not Path(protocol_path).is_file() or hashlib.sha256(Path(protocol_path).read_bytes()).hexdigest()!=protocol_sha):
        raise DataContractError("研究方案文件在运行中被修改；请固定方案文件后重新运行")
    extra_csv=[name for name in input_hashes if name not in {t+".csv" for t in TABLE_COLUMNS}]
    followup_cols=[c for c in ["followup_status","followup_status_date"] if c in frames["patients_baseline"].columns]
    followup_missing_centers=[{"project_id":pid,"center_code":cc,"missing_columns":sorted(c for c in cols.get("patients_baseline",[]) if c.startswith("followup_status"))}
                              for (pid,cc),cols in sorted(source_missing.items()) if any(c.startswith("followup_status") for c in cols.get("patients_baseline",[]))]
    provenance=provenance_files(data_dir,input_hashes)
    log={"additional_csv_not_analyzed":extra_csv,"analysis_version":VERSION,"generated_at":datetime.now(timezone.utc).isoformat(),
         "mode":"protocol" if protocol else "candidates_only",
         "status":"completed_protocol_review" if protocol else "completed_descriptive_review",
         "patients_n":len(patients),"visits_n":len(visits),"ols_evaluable_n":int((slopes.status=="descriptive_ols").sum()),"annual_selected_n":int((outcomes.status=="visit_selected").sum()),"qc_issue_count":len(qc),"input_dir":str(data_dir),"input_sha256":input_hashes,"script_sha256":sha256(__file__),
         "environment":{"python":platform.python_version(),"python_implementation":platform.python_implementation(),"python_executable_version":sys.version.split()[0],"platform":platform.platform()},
         "dependencies":{x:importlib.metadata.version(x) for x in ["pandas","numpy","openpyxl","matplotlib"]},
         "data_dictionary_sha256":(provenance["data_dictionary.json"] or {}).get("sha256"),
         "provenance_files":provenance,"merge_manifest":merge_info,
         "identity_key":KEY,"missingness":"no imputation; see per-record status","statsmodels_used":False,"kfre_used":False,"endpoint_adjudication":False,
         "overrides":{"allow_fully_excluded_centers":bool(allow_fully_excluded_centers),"allow_unmatched_protocol_terms":bool(allow_unmatched_protocol_terms)},
         "consent_filter":consent_info,
         "egfr_sources":egfr_source_summary(visits),
         "followup_status":{"columns_present":followup_cols,"used_for_censoring":bool(protocol and protocol["censoring"]["use_followup_status"] and "followup_status" in followup_cols),
                            "centers_without_followup_status_in_source_export":followup_missing_centers,
                            "note":("" if followup_cols else "未提供随访状态列（followup_status / followup_status_date）")
                                   +("；合并输入中以下中心的源导出缺少随访状态列（以空值填充），这些患者未按随访状态截尾："+"、".join(f"{c['project_id']}/{c['center_code']}" for c in followup_missing_centers) if followup_missing_centers else "")}}
    if protocol:
        log["protocol_endpoint_determination"]=True
        status_counts={}
        for (module,endpoint,status),n in (endpoints.groupby(["module","endpoint","status"]).size().items() if len(endpoints) else []):
            status_counts.setdefault(module,{})[f"{endpoint}|{status}"]=int(n)
        log["protocol"]={"tag":psummary["tag"],"protocol_id":protocol["protocol_id"],"version":protocol["version"],"file_name":Path(protocol_path).name,
                         "sha256":protocol_sha,"confirmed_by":protocol["confirmation"]["confirmed_by"],"confirmed_date":protocol["confirmation"]["confirmed_date"],
                         "sections_computed":psummary["sections_computed"],"sections_skipped":psummary["sections_skipped"],
                         "modules_not_covered_for_remission":psummary["modules_not_covered_for_remission"],
                         "egfr_decline_kidney_failure_or_death_before_confirmation":psummary["egfr_decline_kidney_failure_or_death_before_confirmation"],
                         "term_matches":term_matches,
                         "visit_upcr_unit_declarations":unit_declarations,"units":protocol["units"],
                         "endpoint_rows_n":len(endpoints),"endpoint_status_counts_by_module":status_counts,
                         "text":protocol_raw_text}
    else:
        log["protocol_endpoint_determination"]=False
        log["protocol"]=None
        log["candidate_note"]="未提供研究方案（--protocol）：仅输出未确认的单次数值候选，不是终点，也未汇总为比例"
    output_dir.parent.mkdir(parents=True,exist_ok=True)
    stage=Path(tempfile.mkdtemp(prefix=".analysis-stage-",dir=output_dir.parent))
    try:
        tables={"patients_derived.csv":patients,"visits_derived.csv":visits,"labs_review.csv":labs,"outcomes_12m.csv":outcomes,"manual_events_review.csv":events,"egfr_slope_per_patient.csv":slopes,"qc_issues.csv":qc}
        if protocol:
            tables["endpoints_per_protocol.csv"]=endpoints
            tables["timepoints_per_protocol.csv"]=timepoints
        else:
            tables["candidates_unconfirmed_egfr_decline.csv"]=declines
            tables["candidates_unconfirmed_igan_proteinuria.csv"]=protein
        for metric in ["egfr","upcr"]:
            trend=trend_summary(visits,metric)
            tables[f"trend_{metric}_patient_month.csv"]=trend
            plot_descriptive_trend(trend,metric,stage/f"plot_{metric}_trend.png")
        for name,df in tables.items(): df.to_csv(stage/name,index=False,encoding="utf-8-sig",date_format="%Y-%m-%d")
        write_xlsx(stage/"table1_baseline.xlsx",{"Table1":table_one(patients,visits,frames["meds_long"],consent_info)})
        write_xlsx(stage/"qc_report.xlsx",{"Issues":qc})
        if protocol:
            # Exactly the bytes that were validated and hashed above; the file is not read again.
            (stage/"protocol_used.json").write_bytes(protocol_bytes)
            adjudication={"status":"protocol_parameterized_determination","protocol":psummary["tag"],"protocol_sha256":protocol_sha,"manual_clinical_adjudication":False,"sections_computed":psummary["sections_computed"]}
        else:
            adjudication={"status":"not_executed_unconfirmed_candidates_only","note":"supply a PI-confirmed --protocol file to determine sustained endpoints"}
        models={"KFRE":{"status":"disabled_pending_validated_model_and_UACR_contract","risk_values_emitted":False},"LME":{"status":"not_executed_no_protocol_specification"},"clinical_endpoint_adjudication":adjudication}
        (stage/"MODEL_STATUS.json").write_text(json.dumps(models,indent=2,ensure_ascii=False),encoding="utf-8")
        (stage/"METHODS_ACTUAL_EN.md").write_text(actual_methods(log,protocol_raw_text),encoding="utf-8")
        log["output_sha256"]={p.name:sha256(p) for p in sorted(stage.iterdir())}
        log["generated_files"]=sorted(log["output_sha256"])+["RUN_LOG.json"]
        (stage/"RUN_LOG.json").write_text(json.dumps(log,indent=2,ensure_ascii=False,allow_nan=False),encoding="utf-8")
        if output_dir.exists():
            keep=output_dir/".keep"
            if keep.exists(): keep.unlink()
            output_dir.rmdir()
        os.replace(stage,output_dir)
    finally:
        if stage.exists(): shutil.rmtree(stage)
    return log


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data",type=Path,default=BASE/"data")
    parser.add_argument("--out",type=Path,default=None,help="New/empty directory; default outputs/run-UTC timestamp")
    parser.add_argument("--protocol",type=Path,default=None,help="PI-confirmed study protocol JSON (see protocol_template.json). Without it only unconfirmed candidates are produced.")
    parser.add_argument("--allow-fully-excluded-centers",action="store_true",help="Proceed although a project/center (or the whole input) has no patient with consent_status consented/waiver. Recorded in RUN_LOG.json.")
    parser.add_argument("--allow-unmatched-protocol-terms",action="store_true",help="Proceed although a confirmed protocol module or censoring event type matches no record. Recorded in RUN_LOG.json.")
    args=parser.parse_args()
    out=args.out or BASE/"outputs"/("run-"+datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ"))
    try:
        log=analyze(args.data,out,args.protocol,allow_fully_excluded_centers=args.allow_fully_excluded_centers,allow_unmatched_protocol_terms=args.allow_unmatched_protocol_terms)
    except (DataContractError,OSError,ValueError) as exc:
        parser.exit(2,f"ANALYSIS_FAILED: {exc}\nNo successful report produced; resolve source issues then use a new run directory.\n")
    print(f"DESCRIPTIVE_REVIEW_OK: {out.resolve()}\nPatients: {log['patients_n']}; visits: {log['visits_n']}; QC issues: {log['qc_issue_count']}. Review MODEL_STATUS and METHODS_ACTUAL_EN before using results.")
    consent=log["consent_filter"]
    if consent["column_present"]:
        print(f"Consent filter: excluded {consent['excluded_n']} patients without consented/waiver status; per project/center counts in RUN_LOG.json and qc_issues.csv.")
        if {"unconfirmed","(empty)"} & set(consent["excluded_by_status"]):
            print("知情同意提示："+CONSENT_HINT_ZH+"\nConsent note: "+CONSENT_HINT_EN)
        if consent["fully_excluded_centers"]:
            print("WARNING: project/center groups with no included patient (allowed by --allow-fully-excluded-centers): "+", ".join(consent["fully_excluded_centers"]))
    else:
        print("Consent filter: no consent_status column; not applied.")
    if log["followup_status"]["centers_without_followup_status_in_source_export"]:
        print("WARNING: "+log["followup_status"]["note"])
    if log["protocol"]:
        p=log["protocol"]
        print(f"Protocol {p['tag']} (sha256 {p['sha256'][:12]}…): computed {', '.join(p['sections_computed']) or 'no section (none PI-confirmed)'}.")
        if p["term_matches"] and p["term_matches"]["unmatched"]:
            print("WARNING: protocol terms with zero matches (allowed by --allow-unmatched-protocol-terms): "+"; ".join(p["term_matches"]["unmatched"]))
    else:
        print("No --protocol: candidates are UNCONFIRMED single-measurement flags, not endpoints.")


if __name__=="__main__":
    main()
